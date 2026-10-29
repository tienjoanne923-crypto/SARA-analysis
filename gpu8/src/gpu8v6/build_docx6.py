"""Render the V6 manuscript Markdown to DOCX with real Word equations (section 17.3).

Display equations written as ``$$ ... $$`` are converted LaTeX -> MathML
(latex2mathml, installed outside the frozen experiment environment in
``gpu8/tools_libs``) -> OMML (Office's own MML2OMML.XSL) and inserted as native
``m:oMath`` objects, not as centred plain text.  Text, tables, figures and the
SimSun / black / three-line-table contract are inherited from the V5 builder.
After saving, the document is reopened and the equation objects are counted.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

from gpu8 import paths as p5
from gpu8.build_docx import (BODY_FONT, BODY_PT, LATIN_FONT, add_paragraph, add_table,
                             parse_markdown_table)

from . import paths6
from .common import dump_json

sys.path.insert(0, str(p5.GPU8_ROOT / "tools_libs"))
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")


class EquationConverter:
    def __init__(self):
        import latex2mathml.converter as l2m
        from lxml import etree
        self.l2m = l2m
        self.etree = etree
        self.xslt = etree.XSLT(etree.parse(str(MML2OMML)))

    def omml(self, latex: str) -> str:
        mathml = self.l2m.convert(latex)
        # aligned environments can leave bare ampersands in the MathML text nodes
        mathml = re.sub(r"&(?![a-zA-Z#][a-zA-Z0-9]*;)", "&amp;", mathml)
        root = self.etree.fromstring(mathml.encode("utf-8"))
        out = self.xslt(root).getroot()
        return self.etree.tostring(out, encoding="unicode")


MATH_RE = re.compile(r"(\$[^$]+\$)")


def add_rich(p, text: str, conv: "EquationConverter", size=BODY_PT, bold=False, stats=None):
    """Runs for text (with **bold** / *italic*) and inline OMML for $...$ segments."""
    from gpu8.build_docx import parse_inline, set_run
    for part in MATH_RE.split(text or ""):
        if not part:
            continue
        if part.startswith("$") and part.endswith("$") and len(part) > 2:
            try:
                p._p.append(parse_xml(conv.omml(part[1:-1])))
                if stats is not None:
                    stats["inline_equations"] += 1
                continue
            except Exception as exc:
                if stats is not None:
                    stats["equation_failures"].append({"latex": part, "error": str(exc)})
        for seg, b, it in parse_inline(part):
            set_run(p.add_run(seg), size=size, bold=bold or b, italic=it)
    return p


def add_par_m(doc, text, conv, stats, size=BODY_PT, bold=False, align=None, space_after=6,
              first_line_indent=None, left_indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    if align is not None:
        p.alignment = align
    if first_line_indent:
        p.paragraph_format.first_line_indent = Pt(first_line_indent)
    if left_indent:
        p.paragraph_format.left_indent = Pt(left_indent)
    return add_rich(p, text, conv, size=size, bold=bold, stats=stats)


def add_table_m(doc, rows, conv, stats):
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from gpu8.build_docx import set_cell_border
    if not rows:
        return None
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
            add_rich(p, row[j] if j < len(row) else "", conv, size=9, bold=(i == 0), stats=stats)
            set_cell_border(cell, top=(i == 0), bottom=(i in (0, len(rows) - 1)))
    return t


def build(md_path: Path, out_path: Path, figure_dir: Path | None = None) -> dict:
    conv = EquationConverter()
    text = Path(md_path).read_text(encoding="utf-8").lstrip("﻿")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.1)
        section.right_margin = Inches(1.1)
    style = doc.styles["Normal"]
    style.font.name = LATIN_FONT
    style.font.size = Pt(BODY_PT)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    lines = text.split("\n")
    i = 0
    stats = {"paragraphs": 0, "tables": 0, "images": 0, "headings": 0, "equations": 0,
             "inline_equations": 0, "equation_failures": [], "missing_figures": []}
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and \
                re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table_m(doc, parse_markdown_table(block), conv, stats)
            stats["tables"] += 1
            add_paragraph(doc, "", space_after=4)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level, title = len(m.group(1)), m.group(2)
            size = {1: 15, 2: 13, 3: 11.5, 4: 10.5}[level]
            p = add_par_m(doc, title, conv, stats, size=size, bold=True, space_after=8)
            p.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
            stats["headings"] += 1
            i += 1
            continue
        if stripped.startswith("$$"):
            block = [stripped]
            if not (stripped.endswith("$$") and len(stripped) > 2):
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    block.append(lines[i])
                    i += 1
                if i < len(lines):
                    block.append(lines[i])
            i += 1
            tex = " ".join(block).replace("$$", "").strip()
            label = ""
            mm = re.search(r"\\tag\{([^}]*)\}", tex)
            if mm:
                label = mm.group(1)
                tex = tex.replace(mm.group(0), "").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                p._p.append(parse_xml(conv.omml(tex)))
                stats["equations"] += 1
            except Exception as exc:  # recorded, rendered as text so nothing is lost
                stats["equation_failures"].append({"latex": tex, "error": str(exc)})
                p.add_run(tex)
            if label:
                r = p.add_run(f"    ({label})")
                r.font.size = Pt(10)
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if m:
            rel = m.group(2)
            src = (figure_dir / Path(rel).name) if figure_dir else Path(rel)
            if src.exists():
                doc.add_picture(str(src), width=Inches(6.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                stats["images"] += 1
            else:
                stats["missing_figures"].append(str(src))
                add_paragraph(doc, f"[图缺失：{src.name}]", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            i += 1
            continue
        if re.match(r"^[-*]\s+", stripped) or re.match(r"^\d+\.\s+", stripped):
            body = re.sub(r"^([-*])\s+", "", stripped)
            add_par_m(doc, body, conv, stats, space_after=3, left_indent=18)
            stats["paragraphs"] += 1
            i += 1
            continue
        if stripped.startswith(">"):
            add_par_m(doc, stripped.lstrip("> ").strip(), conv, stats, space_after=6, left_indent=24)
            stats["paragraphs"] += 1
            i += 1
            continue
        add_par_m(doc, stripped, conv, stats, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=21)
        stats["paragraphs"] += 1
        i += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    with zipfile.ZipFile(out_path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    stats["verified_omml_objects"] = xml.count("<m:oMath>") + xml.count("<m:oMath ")
    check = Document(str(out_path))
    stats["verified_paragraphs"] = len(check.paragraphs)
    stats["verified_tables"] = len(check.tables)
    stats["bytes"] = out_path.stat().st_size
    stats["path"] = str(out_path)
    return stats


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--markdown", required=True)
    ap.add_argument("--out", default="")
    ap.add_argument("--figure_dir", default="")
    a = ap.parse_args(argv)
    md = Path(a.markdown)
    out = Path(a.out) if a.out else md.with_suffix(".docx")
    fig = Path(a.figure_dir) if a.figure_dir else paths6.FIGURES
    st = build(md, out, fig)
    dump_json(paths6.REPORTS / "docx_build_report_v6.json", st)
    print(json.dumps(st, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
