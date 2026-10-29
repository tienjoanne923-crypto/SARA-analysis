"""MathType equation factory.

LaTeX -> OMML (Word's own MML2OMML.XSL via build_docx6.EquationConverter) in a scratch document, one equation per
paragraph (inline equations are written between two marker characters so that Word keeps them inline; display
equations stand alone).  A hidden Word instance then runs MathType's own "Convert Equations" (OMML -> MathType OLE),
driven by UI Automation (tools/mathtype/convert_docx_to_mathtype.ps1 + dialog_controller.ps1).  Every converted
object is harvested per paragraph: the OLE storage (Equation.DSMT4, with the MTEF made by MathType itself), the
WMF preview, its size and baseline offset.  The final documents are assembled from these assets by build_docx10.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from lxml import etree

from .numbers9 import ROOT

sys.path.append(r"D:\信\gpu8\tools_libs10")             # olefile (appended: the interpreter's packages win)
CACHE = ROOT / "build" / "mathtype"
TOOLS = Path(r"D:\信\gpu8\submission9\tools\mathtype")
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "v": "urn:schemas-microsoft-com:vml", "o": "urn:schemas-microsoft-com:office:office",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}


# LaTeX -> OMML -> MathType loses two styles: \mathbf becomes a Unicode math-bold letter that MathType maps back to an
# italic variable, and a lone upright letter from \mathrm is merged into an italic run.  Before conversion \mathbf is
# written as \boldsymbol (bold italic, the ISO 80000-2 style for vectors and matrices; MathType keeps it as an explicit
# bold-italic font style) and \mathrm as \text (MathType's upright Text style).  PREP versions the cache.
PREP = "p2"


def prepare(tex: str) -> str:
    return tex.strip().replace("\\mathbf{", "\\boldsymbol{").replace("\\mathrm{", "\\text{")


def eq_key(tex: str, mode: str) -> str:
    return hashlib.sha1(f"{PREP}|{mode}|{tex.strip()}".encode("utf-8")).hexdigest()[:16]


def make_scratch(jobs: list[tuple[str, str, str]], path: Path) -> None:
    """jobs = [(key, tex, mode)], mode in {inline, display}."""
    from docx import Document
    from docx.oxml import parse_xml
    from gpu8v6.build_docx6 import EquationConverter
    conv = EquationConverter()
    doc = Document()
    for key, tex, mode in jobs:
        p = doc.add_paragraph()
        if mode == "inline":
            p.add_run("\u00b7")
        p._p.append(parse_xml(conv.omml(prepare(tex))))
        if mode == "inline":
            p.add_run("\u00b7")
    doc.save(str(path))


def convert(scratch: Path, out: Path, timeout: int = 1200) -> str:
    if out.exists():
        out.unlink()
    cmd = ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(TOOLS / "convert_docx_to_mathtype.ps1"),
           "-InputPath", str(scratch), "-OutputPath", str(out), "-TimeoutSeconds", str(timeout)]
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout + 120)
    log = r.stdout + "\n" + r.stderr
    if r.returncode != 0:
        raise RuntimeError(f"MathType conversion failed ({r.returncode}):\n{log[-3000:]}")
    return log


def _style_pt(style: str, key: str) -> float:
    m = re.search(rf"{key}:([\d.]+)pt", style)
    return float(m.group(1)) if m else float("nan")


def audit_bin(blob: bytes) -> dict:
    import olefile
    ole = olefile.OleFileIO(io.BytesIO(blob))
    try:
        streams = ["/".join(s) for s in ole.listdir()]
        native = ole.openstream("Equation Native").read()
        clsid = ole.root.clsid
    finally:
        ole.close()
    hdr_len = int.from_bytes(native[:2], "little")
    cb = int.from_bytes(native[8:12], "little")
    mtef = native[hdr_len:hdr_len + cb]
    return {"streams": streams, "clsid": clsid, "hdr_len": hdr_len, "cbObject": cb, "mtef_len": len(mtef),
            "mtef_head": mtef[:6].hex(), "mtef_ok": mtef[:2] == b"\x05\x01" and cb == len(native) - hdr_len,
            "replacement_chars": mtef.count(b"\xfd\xff")}


def harvest(converted: Path, jobs: list[tuple[str, str, str]]) -> dict:
    """one converted object per scratch paragraph, in order; assets are cached by key."""
    CACHE.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(converted) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        rels = etree.fromstring(z.read("word/_rels/document.xml.rels"))
        target = {r.get("Id"): r.get("Target") for r in rels.findall("rel:Relationship", NS)}
        paras = root.findall(".//w:body/w:p", NS)
        if len(root.findall(".//m:oMath", NS)):
            raise RuntimeError("OMML left after conversion")
        report = {}
        for (key, tex, mode), p in zip(jobs, paras):
            objs = p.findall(".//w:object", NS)
            if len(objs) != 1:
                raise RuntimeError(f"paragraph for {key} has {len(objs)} objects")
            ob = objs[0]
            run = ob.getparent()
            pos = run.find("w:rPr/w:position", NS)
            shape = ob.find("v:shape", NS)
            img = shape.find("v:imagedata", NS)
            ole = ob.find("o:OLEObject", NS)
            style = shape.get("style", "")
            wmf_name = target[img.get(f"{{{NS['r']}}}id")]
            bin_name = target[ole.get(f"{{{NS['r']}}}id")]
            blob, pic = z.read("word/" + bin_name), z.read("word/" + wmf_name)
            ext = Path(wmf_name).suffix.lower()
            (CACHE / f"{key}.bin").write_bytes(blob)
            (CACHE / f"{key}{ext}").write_bytes(pic)
            info = {"key": key, "tex": tex, "mode": mode, "progid": ole.get("ProgID"), "image_ext": ext,
                    "dxa": int(ob.get(f"{{{NS['w']}}}dxaOrig")), "dya": int(ob.get(f"{{{NS['w']}}}dyaOrig")),
                    "width_pt": _style_pt(style, "width"), "height_pt": _style_pt(style, "height"),
                    "position": int(pos.get(f"{{{NS['w']}}}val")) if pos is not None else 0,
                    **audit_bin(blob)}
            (CACHE / f"{key}.json").write_text(json.dumps(info, ensure_ascii=False, indent=1), encoding="utf-8")
            report[key] = info
    return report


# ------------------------------------------------------------------ equation size (tables 9 pt, legends 10.5 pt)
# MathType converts with its 12 pt preferences.  An equation set in smaller text is given the same MTEF with the
# absolute sizes of its EQN_PREFS record (full/script/symbol sizes and the point-valued spacings) scaled, so that
# MathType re-renders it at that size when it is edited; the stored 12 pt preview is displayed at the scaled size.
_DIGITS = "0123456789.-"
_STREAM_ORDER = ["\x01CompObj", "\x01Ole", "\x03ObjInfo", "Equation Native"]


def _dims_decode(buf: bytes, j: int, count: int) -> tuple[list, int]:
    vals, unit, cur = [], None, ""
    while len(vals) < count:
        b = buf[j]
        j += 1
        for n in (b >> 4, b & 15):
            if len(vals) >= count:
                break                                   # pad nibble of the last byte
            if unit is None:
                unit, cur = n, ""
            elif n == 15:
                vals.append((unit, cur))
                unit = None
            else:
                cur += _DIGITS[n]
    return vals, j


def _dims_encode(vals: list) -> bytes:
    nib = []
    for unit, v in vals:
        nib += [unit] + [_DIGITS.index(c) for c in v] + [15]
    if len(nib) % 2:
        nib.append(0)
    return bytes((nib[k] << 4) | nib[k + 1] for k in range(0, len(nib), 2))


def _num(x: float) -> str:
    s = f"{x:.2f}".rstrip("0").rstrip(".")
    return s or "0"


def _prefs_span(mtef: bytes) -> tuple[int, int]:
    """(start, end) of the EQN_PREFS record in an MTEF v5 stream."""
    if mtef[0] != 5:
        raise ValueError("not MTEF v5")
    i = mtef.index(b"\x00", 5) + 2                     # application key, then the equation-options byte
    while True:
        tag = mtef[i]
        if tag == 19:                                   # ENCODING_DEF
            i = mtef.index(b"\x00", i + 1) + 1
        elif tag == 17:                                 # FONT_DEF
            i = mtef.index(b"\x00", i + 2) + 1
        elif tag == 18:                                 # EQN_PREFS
            j = i + 2
            for _ in range(2):
                _, j = _dims_decode(mtef, j + 1, mtef[j])
            n = mtef[j]
            j += 1
            for _ in range(n):
                j += 2 if mtef[j] else 1
            return i, j
        else:
            raise ValueError(f"unexpected MTEF record {tag} before EQN_PREFS")


def scale_mtef(mtef: bytes, factor: float) -> bytes:
    i, end = _prefs_span(mtef)
    out = bytearray(mtef[i:i + 2])
    j = i + 2
    for _ in range(2):
        cnt = mtef[j]
        vals, j = _dims_decode(mtef, j + 1, cnt)
        vals = [(u, _num(float(v) * factor)) if u in (0, 1, 2, 3) else (u, v) for u, v in vals]
        out += bytes([cnt]) + _dims_encode(vals)
    out += mtef[j:end]
    return mtef[:i] + bytes(out) + mtef[end:]


def prefs_sizes(mtef: bytes) -> list:
    i, _ = _prefs_span(mtef)
    return _dims_decode(mtef, i + 3, mtef[i + 2])[0]


def _write_storage(streams: dict, clsid: str) -> bytes:
    import os
    import tempfile
    import pythoncom
    import pywintypes
    from win32com import storagecon as sc
    (CACHE / "work").mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".stg", dir=CACHE / "work")
    os.close(fd)
    os.unlink(tmp)
    stg = pythoncom.StgCreateDocfile(tmp, sc.STGM_CREATE | sc.STGM_READWRITE | sc.STGM_SHARE_EXCLUSIVE, 0)
    stg.SetClass(pywintypes.IID("{" + clsid + "}"))
    for name in _STREAM_ORDER:
        st = stg.CreateStream(name, sc.STGM_CREATE | sc.STGM_WRITE | sc.STGM_SHARE_EXCLUSIVE, 0, 0)
        st.Write(streams[name])
        st = None
    stg.Commit(0)
    stg = None
    data = Path(tmp).read_bytes()
    os.unlink(tmp)
    return data


def scale_bin(blob: bytes, factor: float) -> bytes:
    import olefile
    ole = olefile.OleFileIO(io.BytesIO(blob))
    try:
        names = ["/".join(s) for s in ole.listdir()]
        assert sorted(names) == sorted(_STREAM_ORDER), names
        streams = {n: ole.openstream(n).read() for n in names}
        clsid = ole.root.clsid
    finally:
        ole.close()
    native = streams["Equation Native"]
    hdr_len = int.from_bytes(native[:2], "little")
    mtef = scale_mtef(native[hdr_len:hdr_len + int.from_bytes(native[8:12], "little")], factor)
    hdr = bytearray(native[:hdr_len])
    hdr[8:12] = len(mtef).to_bytes(4, "little")
    streams["Equation Native"] = bytes(hdr) + mtef
    return _write_storage(streams, clsid)


def sized(info: dict, factor: float) -> dict:
    """asset of `info` for text `factor` x 12 pt (cached as <key>_<per mille>)."""
    if abs(factor - 1) < 1e-9:
        return {**info, "bin": str(CACHE / f"{info['key']}.bin"), "image": str(CACHE / f"{info['key']}{info['image_ext']}"),
                "factor": 1.0}
    tag = f"{info['key']}_{round(factor * 1000)}"
    rt = CACHE / f"{tag}_rt.json"                          # re-typeset by MathType at this size (mt_resize10)
    if rt.exists():
        out = json.loads(rt.read_text(encoding="utf-8"))
        return {**out, "bin": str(CACHE / f"{tag}_rt.bin"), "image": str(CACHE / f"{tag}_rt{out['image_ext']}"),
                "factor": factor, "retypeset": True}
    j = CACHE / f"{tag}.json"
    if j.exists():
        return json.loads(j.read_text(encoding="utf-8"))
    blob = scale_bin((CACHE / f"{info['key']}.bin").read_bytes(), factor)
    (CACHE / f"{tag}.bin").write_bytes(blob)
    out = {**info, "key": tag, "base_key": info["key"], "factor": factor,
           "dxa": round(info["dxa"] * factor), "dya": round(info["dya"] * factor),
           "width_pt": round(info["width_pt"] * factor, 2), "height_pt": round(info["height_pt"] * factor, 2),
           "position": round(info["position"] * factor), "bin": str(CACHE / f"{tag}.bin"),
           "image": str(CACHE / f"{info['key']}{info['image_ext']}"), **audit_bin(blob)}
    out["prefs_sizes"] = prefs_sizes(_mtef_of(blob))
    j.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    return out


def _mtef_of(blob: bytes) -> bytes:
    import olefile
    ole = olefile.OleFileIO(io.BytesIO(blob))
    try:
        native = ole.openstream("Equation Native").read()
    finally:
        ole.close()
    h = int.from_bytes(native[:2], "little")
    return native[h:h + int.from_bytes(native[8:12], "little")]


# ------------------------------------------------------------------ collection
DISPLAY_RX = re.compile(r"^\$\$(.+?)\$\$\s*$")
INLINE_RX = re.compile(r"\$([^$]+)\$")


def collect(md_paths: list[Path]) -> list[tuple[str, str]]:
    """unique (tex, mode) of every display ($$...$$ line) and inline ($...$) equation, in order of appearance."""
    seen, out = set(), []
    for p in md_paths:
        for line in Path(p).read_text(encoding="utf-8").split("\n"):
            s = line.strip()
            m = DISPLAY_RX.match(s)
            items = [(m.group(1).strip(), "display")] if m else [(t.strip(), "inline") for t in INLINE_RX.findall(s)]
            for it in items:
                if it not in seen:
                    seen.add(it)
                    out.append(it)
    return out


def build_assets(eqs: list[tuple[str, str]], tag: str) -> dict:
    """eqs = [(tex, mode)]; converts the ones not yet cached and returns {key: info} for all."""
    jobs, have = [], {}
    for tex, mode in eqs:
        k = eq_key(tex, mode)
        j = CACHE / f"{k}.json"
        if j.exists():
            have[k] = json.loads(j.read_text(encoding="utf-8"))
        elif all(k != x[0] for x in jobs):
            jobs.append((k, tex.strip(), mode))
    if jobs:
        work = CACHE / "work"
        work.mkdir(parents=True, exist_ok=True)
        scratch, out = work / f"scratch_{tag}.docx", work / f"scratch_{tag}_mt.docx"
        make_scratch(jobs, scratch)
        log = convert(scratch, out)
        (work / f"scratch_{tag}.log").write_text(log, encoding="utf-8")
        have.update(harvest(out, jobs))
    return have


def build_all(md_paths: list[Path], tag: str = "all") -> dict:
    eqs = collect(md_paths)
    rep = build_assets(eqs, tag)
    bad = {k: v["tex"] for k, v in rep.items() if not v["mtef_ok"] or v["replacement_chars"] or v["progid"] != "Equation.DSMT4"}
    summary = {"equations": len(eqs), "display": sum(1 for _, m in eqs if m == "display"), "assets": len(rep), "bad": bad,
               "widest_display_pt": max((v["width_pt"] for v in rep.values() if v["mode"] == "display"), default=0)}
    (CACHE / f"summary_{tag}.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8")
    return summary


if __name__ == "__main__" and len(sys.argv) > 1:
    print(json.dumps(build_all([Path(a) for a in sys.argv[1:]]), ensure_ascii=False, indent=1))
elif __name__ == "__main__":
    tests = [(r"x^{2}+\frac{a}{b}", "inline"), (r"s_F", "inline"), (r"\mathcal{L}_{DG}", "inline"),
             (r"T_j\setminus I_{jb}", "inline"), (r"\lVert q_d\odot \mathrm{IC}\rVert", "inline"),
             (r"s_d(g)=\sum_{d'\in\mathcal{N}_{20}(d)}\frac{\mathrm{sim}(d,d')}{\sum_{d''}\mathrm{sim}(d,d'')}\,\frac{A_{d'g}}{\deg(d')}", "display"),
             (r"z_d^{\top}z^F_g", "inline")]
    rep = build_assets(tests, "probe")
    print(json.dumps({k: {kk: v[kk] for kk in ("mode", "progid", "width_pt", "height_pt", "position", "mtef_ok",
                                               "replacement_chars", "image_ext")} for k, v in rep.items()}, indent=1))
