"""Checks of the SARA submission package; exits non-zero when any check fails.

  1 forbidden (internal / process) wording in the main text, supplement, Word files, figure text and public CSVs;
  2 number tracing: every decimal, percentage and thousands-separated integer of the text was produced by fmt9
    from the saved results (or is a listed design constant);
  3 formatting: no signed zero, no "nan", no ASCII hyphen used as a minus sign, no nested full-width brackets;
  4 E4 counts equal ``e4_summary``;  5 references verified, cited in order, no duplicate DOI;
  6 cross references (figures, panels, tables, notes) exist and main figures/tables are cited in order;
  7 Word build report;  8 colour-palette checks of the figures;  9 legend panels match the figure panels.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pandas as pd

from .manuscript9 import STEM
from .numbers9 import FIG, ROOT, e4_summary, load_R9
from .publabel9 import FORBIDDEN_RX
from .supp9 import SUPP_STEM

SUPP_ALLOW = {"非劣"}
NUM_RX = re.compile(r"(?<![\w.])[+−]?\d{1,3}(?:,\d{3})+(?:\.\d+)?(?![\d.])%?|(?<![\w.])[+−]?\d+\.\d+(?![\d.])%?|(?<![\w.,])\d+%")
# design constants and definitions that are part of the method description, not results
CONSTANTS = {"0.05", "0.25", "0.3", "0.5", "0.7", "0.9", "0.8", "0.15", "0.005", "−0.005", "0.02", "0.2", "0.1", "0.001",
             "0.0001", "1.0", "0.03", "12.8", "2.2", "1,500", "10,000", "1,000", "2,000", "0.20", "0.50", "0.30", "0.15",
             "0.25", "10%", "20%", "0.5）", "1.9", "95%", "0.01"}
PANELS = {"1": "abcdefg", "2": "abcde", "3": "abcdefghi", "4": "abcdefg", "5": "abcde", "6": "abcdef", "7": "abcd", "8": "abcde",
          "9": "abcdefg", "10": "ab", "S1": "ab", "S2": "ab", "S3": "abc", "S4": "ab", "S5": "ab", "S6": "ab", "S7": "ab",
          "S8": "abcdefg"}


def _body(md: str) -> str:
    """text without the reference list and without LaTeX."""
    md = re.split(r"\n## (?:补充)?参考文献\n", md)[0] + ("\n" + md.split("## 表\n", 1)[1] if "## 表\n" in md else "")
    md = re.sub(r"\$\$.*?\$\$", " ", md, flags=re.S)
    return re.sub(r"\$[^$]+\$", " ", md)


def check_forbidden(md: str, allow=frozenset()) -> list[str]:
    out = []
    for i, line in enumerate(md.splitlines(), 1):
        for m in FORBIDDEN_RX.finditer(line):
            if not any(a in line[max(0, m.start() - 2):m.end() + 2] for a in allow):
                out.append(f"{i}: {m.group(0)} :: {line[:90]}")
    return out


SCI_RX = re.compile(r"\d+\.\d+ × 10[⁻⁰¹²³⁴⁵⁶⁷⁸⁹]+")


def check_numbers(md: str, trace: dict) -> list[str]:
    """every result number was formatted by fmt9 (section numbers, cross references and constants excepted)."""
    bad = []
    for line in _body(md).splitlines():
        s = line.strip()
        if re.match(r"^\d+\. ", s) or s.startswith("#"):
            continue
        for m in SCI_RX.finditer(line):
            if m.group(0) not in trace:
                bad.append(f"{m.group(0)} :: untraced scientific notation")
        line = SCI_RX.sub(" ", line)
        line = re.sub(r"(方法|正文方法|第) ?\d+\.\d+", " ", line)
        for m in NUM_RX.finditer(line):
            tok = m.group(0)
            if tok in trace or tok in CONSTANTS or tok.lstrip("+") in trace:
                continue
            bad.append(f"{tok} :: {line[max(0, m.start() - 30):m.end() + 20]}")
    return bad


def check_format(md: str) -> list[str]:
    body = _body(md)
    out = []
    for rx, why in ((r"[+−-]0\.0+(?![\d])", "signed zero"), (r"[+−]0%", "signed zero %"), (r"\bnan\b", "nan"),
                    (r"(?<![\w\d.–-])-\d", "ASCII hyphen as minus"), (r"（[^（）]*（[^（）]*）[^（）]*）", "nested brackets"),
                    (r"）\s+[、，。；：与和]", "space after bracket"), (r"[\u4e00-\u9fff][A-Za-z0-9]", "no CJK-Latin space"),
                    (r"[A-Za-z0-9%][\u4e00-\u9fff]", "no Latin-CJK space")):
        for m in re.finditer(rx, body):
            ctx = body[max(0, m.start() - 20):m.end() + 20].replace("\n", " ")
            if why.startswith("no ") and re.search(r"(种子|外层|图|表|说明|第)\s?\d|[A-Za-z]+[\u4e00-\u9fff]", ctx) is None:
                pass
            out.append(f"{why}: {ctx}")
    return out


def xrefs(md: str, supp: bool) -> tuple[list[str], dict]:
    body = _body(md)
    bad, first = [], {}
    for m in re.finditer(r"(?<!视)图 (S?\d+)((?:[a-z](?:–[a-z])?)(?:[、,][a-z](?:–[a-z])?)*)?", body):
        n, letters = m.group(1), m.group(2) or ""
        if n not in PANELS:
            bad.append(f"no figure {n}")
            continue
        first.setdefault(n, m.start())
        for part in re.split(r"[、,]", letters):
            if not part:
                continue
            a, b = (part.split("–") + [part])[:2]
            for c in {chr(x) for x in range(ord(a), ord(b) + 1)}:
                if c not in PANELS[n]:
                    bad.append(f"figure {n} has no panel {c}")
    for m in re.finditer(r"表 (S?)(\d+)", body):
        s, k = m.group(1), int(m.group(2))
        if (s and not 1 <= k <= 10) or (not s and not 1 <= k <= 3):
            bad.append(f"no table {s}{k}")
        first.setdefault(f"T{s}{k}", m.start())
    for m in re.finditer(r"补充说明 S(\d+)", body):
        if not 1 <= int(m.group(1)) <= 6:
            bad.append(f"no note S{m.group(1)}")
    if not supp:
        order = sorted((v, k) for k, v in first.items() if k.isdigit())
        nums = [int(k) for _, k in order]
        if nums != sorted(nums):
            bad.append(f"main figures cited out of order: {nums}")
        torder = [int(k[1:]) for _, k in sorted((v, k) for k, v in first.items() if re.fullmatch(r"T\d+", k))]
        if torder != sorted(torder):
            bad.append(f"main tables cited out of order: {torder}")
    return bad, first


def legend_panels(md: str) -> list[str]:
    bad = []
    for m in re.finditer(r"\*\*图 (S?\d+)　[^*]+\*\* (.+)", md):
        n, leg = m.group(1), m.group(2)
        got = "".join(sorted(set(re.findall(r"\*\*([a-z])\*\*", leg))))
        if got != PANELS[n]:
            bad.append(f"legend of figure {n}: panels {got} != {PANELS[n]}")
    return bad


def refs_ok(info: dict) -> list[str]:
    from .refs9 import REFS
    rep = json.loads((Path(r"D:\信\gpu8\submission9") / "refs_verification.json").read_text(encoding="utf-8"))
    bad = []
    order = info["ref_order"]
    if sorted(order.values()) != list(range(1, len(order) + 1)):
        bad.append("reference numbers not consecutive")
    dois = [REFS[k][1] for k in order if REFS[k][1]]
    if len(dois) != len(set(dois)):
        bad.append("duplicate DOI")
    for k in order:
        st = rep.get(k, {}).get("status")
        if st not in ("ok", "ok_noid", "manual_ok"):
            bad.append(f"reference {k}: {st}")
    return bad


def docx_text(path: Path) -> str:
    from docx import Document
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def run() -> dict:
    R = load_R9()
    main_md = (ROOT / f"{STEM}.md").read_text(encoding="utf-8")
    supp_md = (ROOT / f"{SUPP_STEM}.md").read_text(encoding="utf-8")
    info_m = json.loads((ROOT / "build" / "manuscript9_info.json").read_text(encoding="utf-8"))
    info_s = json.loads((ROOT / "build" / "supp9_info.json").read_text(encoding="utf-8"))
    rep: dict = {}
    rep["1_forbidden_main"] = check_forbidden(main_md)
    rep["1_forbidden_supp"] = check_forbidden(supp_md, SUPP_ALLOW)
    for key, fn in (("main", f"{STEM}.docx"), ("review", f"{STEM.replace('投稿版', '审阅版')}.docx"), ("supp", f"{SUPP_STEM}.docx")):
        p = ROOT / fn
        rep[f"1_forbidden_docx_{key}"] = (check_forbidden(docx_text(p), SUPP_ALLOW if key == "supp" else frozenset())
                                          if p.exists() else ["missing"])
    scan = json.loads((FIG / "text_scan.json").read_text(encoding="utf-8"))
    rep["1_forbidden_figures"] = [f"{k}: {v.get('forbidden')}" for k, v in scan.items() if isinstance(v, dict) and v.get("forbidden")]
    csv_bad = []
    for p in sorted((FIG / "source_data").glob("*.csv")):
        t = p.read_text(encoding="utf-8-sig")
        hits = [m.group(0) for m in FORBIDDEN_RX.finditer(t)]
        if hits:
            csv_bad.append(f"{p.name}: {sorted(set(hits))[:5]}")
    rep["1_forbidden_public_csv"] = csv_bad
    rep["2_untraced_main"] = check_numbers(main_md, info_m["trace"])
    rep["2_untraced_supp"] = check_numbers(supp_md, info_s["trace"])
    rep["3_format_main"] = [x for x in check_format(main_md) if not x.startswith("no ")]
    rep["3_format_supp"] = [x for x in check_format(supp_md) if not x.startswith("no ")]
    rep["3_spacing_notes"] = [x for x in check_format(main_md) + check_format(supp_md) if x.startswith("no ")][:40]
    e4s = e4_summary(R)
    want = f"在 {e4s['n_retrain']} 种整链重训的删除条件中，{e4s['n_inc0']} 种条件下"
    rep["4_e4_counts"] = [] if want in main_md else [f"missing: {want}"]
    rep["5_references_main"] = refs_ok(info_m)
    rep["5_references_supp"] = refs_ok(info_s)
    bad_m, first_m = xrefs(main_md, supp=False)
    bad_s, _ = xrefs(supp_md, supp=True)
    rep["6_xrefs_main"], rep["6_xrefs_supp"] = bad_m, bad_s
    rep["9_legend_panels"] = legend_panels(main_md) + legend_panels(supp_md)
    br = ROOT / "build" / "docx_build_report9.json"
    if br.exists():
        b = json.loads(br.read_text(encoding="utf-8"))
        rep["7_docx"] = [] if b.get("all_ok") else [json.dumps({k: {kk: v[kk] for kk in ("a4", "line_numbers", "page_field",
                                                                                           "omml_objects", "equations",
                                                                                           "inline_equations", "missing_figures",
                                                                                           "equation_failures")}
                                                                   for k, v in b.items() if isinstance(v, dict)}, ensure_ascii=False)]
    else:
        rep["7_docx"] = ["docx build report missing"]
    man = json.loads((FIG / "figures_manifest.json").read_text(encoding="utf-8"))
    pal = man.get("_palette_checks", {})
    rep["8_palette"] = [k for k, v in pal.items() if isinstance(v, dict) and not v.get("pass", True)] if pal else ["no palette checks"]
    abstract = re.search(r"## 摘要\n\n(.+?)\n\n", main_md, flags=re.S).group(1)
    rep["_info"] = {"abstract_chars": len(re.sub(r"\s", "", abstract)), "main_chars": len(main_md), "supp_chars": len(supp_md),
                    "refs_main": info_m["n_refs"], "refs_supp": info_s["n_refs"],
                    "key_points": main_md.split("## 要点", 1)[1].split("## 1 引言")[0].count("\n- "),
                    "first_mentions": {k: v for k, v in sorted(first_m.items(), key=lambda t: t[1])}}
    fails = {k: v for k, v in rep.items() if not k.startswith("_") and not k.startswith("3_spacing") and v}
    rep["ok"] = not fails
    (ROOT / "build" / "verify9_report.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    return rep


if __name__ == "__main__":
    r = run()
    show = {k: (v[:12] if isinstance(v, list) else v) for k, v in r.items() if k != "_info"}
    print(json.dumps(show, ensure_ascii=False, indent=1))
    print(json.dumps({k: v for k, v in r["_info"].items() if k != "first_mentions"}, ensure_ascii=False))
    sys.exit(0 if r["ok"] else 1)
