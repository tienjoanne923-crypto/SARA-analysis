"""Word documents of the SARA submission.

Three documents are built from the Markdown written by ``manuscript9`` and ``supp9``:
  * 投稿版  -- A4, continuous line numbers, page numbers, real heading styles, Key Points as bullets, native Word
              equations (OMML, numbered), three-line tables with captions above, references with hanging indent,
              tables and figure legends at the end; no embedded images (figures are uploaded separately);
  * 审阅版  -- the same text with each figure (300 dpi) embedded above its legend;
  * 补充材料 -- notes, tables and legends, with the supplementary figures embedded.
After saving, every document is reopened and checked (A4, line numbering, page field, headings, tables, OMML
objects, images); the report is written to ``build/docx_build_report9.json``.
"""
from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import paths8

sys.path.insert(0, str(paths8.GPU8_ROOT / "tools_libs"))
from gpu8.build_docx import parse_inline, parse_markdown_table, set_cell_border  # noqa: E402
from gpu8v6.build_docx6 import EquationConverter  # noqa: E402

from .manuscript9 import DATE, STEM  # noqa: E402
from .numbers9 import FIG, ROOT  # noqa: E402
from .supp9 import SUPP_STEM  # noqa: E402

LATIN, EA_BODY, EA_HEAD = "Times New Roman", "SimSun", "SimHei"
BLACK = RGBColor(0, 0, 0)
BODY_PT, SMALL_PT, TABLE_PT = 12, 10.5, 9
REVIEW_STEM = STEM.replace("投稿版", "审阅版")
MATH_RE = re.compile(r"(\$[^$]+\$)")
CAPTION_RX = re.compile(r"^(表 S?\d+|（[a-z]）)")
FIG_RX = re.compile(r"^<!-- figure (\S+) -->$")


# ------------------------------------------------------------------ low-level helpers
def fonts(run, ea=EA_BODY, size=BODY_PT, bold=False, italic=False, mono=False):
    return _fonts(run, ea=ea, size=size, bold=bold, italic=italic, mono=mono)


def _fonts(run, ea=EA_BODY, size=BODY_PT, bold=False, italic=False, mono=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    latin = "Consolas" if mono else LATIN
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(k), latin)
    rf.set(qn("w:eastAsia"), ea)
    return run


def style_fonts(style, ea, size, bold=False):
    style.font.name = LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = BLACK
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(k), LATIN)
    rf.set(qn("w:eastAsia"), ea)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(k)) is not None:
            del rf.attrib[qn(k)]


def rich(p, text, conv, stats, size=BODY_PT, bold=False, ea=EA_BODY):
    """Runs for **bold** / *italic* / `code`, inline OMML for $...$ (Markdown escapes such as \\* are unescaped)."""
    esc = ""
    text = (text or "").replace("\\*", esc)

    def fonts(run, **kw):                                   # restore escaped asterisks in the visible text
        run.text = run.text.replace(esc, "*")
        return _fonts(run, **kw)
    for part in MATH_RE.split(text or ""):
        if not part:
            continue
        if part.startswith("$") and part.endswith("$") and len(part) > 2:
            try:
                p._p.append(parse_xml(conv.omml(part[1:-1])))
                stats["inline_equations"] += 1
                continue
            except Exception as exc:  # recorded; the LaTeX text is kept so nothing is lost
                stats["equation_failures"].append({"latex": part, "error": str(exc)})
        pat = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+?`)")
        for seg in pat.split(part):
            if not seg:
                continue
            if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
                fonts(p.add_run(seg[2:-2]), ea=ea, size=size, bold=True)
            elif seg.startswith("`") and seg.endswith("`"):
                fonts(p.add_run(seg[1:-1]), ea=ea, size=size - 1, bold=bold, mono=True)
            elif seg.startswith("*") and seg.endswith("*") and len(seg) > 2:
                fonts(p.add_run(seg[1:-1]), ea=ea, size=size, bold=bold, italic=True)
            else:
                fonts(p.add_run(seg), ea=ea, size=size, bold=bold)
    return p


def para(doc, text, conv, stats, size=BODY_PT, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True,
         before=0, after=4, spacing=1.5, style=None, ea=EA_BODY, keep_next=False, hanging=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), spacing
    if align is not None:
        p.alignment = align
    if indent:
        pf.first_line_indent = Pt(2 * size)
    if hanging:
        pf.left_indent = Pt(hanging)
        pf.first_line_indent = Pt(-hanging)
    pf.keep_with_next = keep_next
    rich(p, text, conv, stats, size=size, bold=bold, ea=ea)
    stats["paragraphs"] += 1
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_field(run, instr: str):
    for tag, extra in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None),
                       ("w:fldChar", {"w:fldCharType": "separate"}), ("w:t", None), ("w:fldChar", {"w:fldCharType": "end"})):
        el = OxmlElement(tag)
        if extra:
            for k, v in extra.items():
                el.set(qn(k), v)
        if tag == "w:instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        if tag == "w:t":
            el.text = "1"
        run._r.append(el)


def setup(doc, line_numbers=True):
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin = s.bottom_margin = Cm(2.54)
        s.left_margin = s.right_margin = Cm(2.54)
        if line_numbers:
            ln = OxmlElement("w:lnNumType")
            ln.set(qn("w:countBy"), "1")
            ln.set(qn("w:restart"), "continuous")
            ln.set(qn("w:distance"), "283")
            s._sectPr.append(ln)
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(fonts(fp.add_run(), size=10), "PAGE")
    st = doc.styles
    style_fonts(st["Normal"], EA_BODY, BODY_PT)
    for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        style_fonts(st[name], EA_HEAD, size, bold=True)
        pf = st[name].paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next = Pt(12 if name != "Heading 3" else 8), Pt(6), True
        pf.line_spacing = 1.3
    st["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # the default Title style draws a coloured bottom border; remove it
    ppr = st["Title"].element.get_or_add_pPr()
    for b in ppr.findall(qn("w:pBdr")):
        ppr.remove(b)
    style_fonts(st["List Bullet"], EA_BODY, BODY_PT)


def table(doc, rows, conv, stats):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
            rich(p, row[j] if j < len(row) else "", conv, stats, size=TABLE_PT, bold=(i == 0))
            set_cell_border(cell, top=(i == 0), bottom=(i in (0, len(rows) - 1)))
    stats["tables"] += 1
    return t


def review_png(stem: str) -> Path:
    """300 dpi copy of the 600 dpi figure (review documents only)."""
    from PIL import Image
    out = ROOT / "build" / "png300" / f"{stem}.png"
    src = FIG / f"{stem}.png"
    if not out.exists() or out.stat().st_mtime < src.stat().st_mtime:
        out.parent.mkdir(parents=True, exist_ok=True)
        im = Image.open(src)
        im = im.convert("RGB")
        im = im.resize((im.width // 2, im.height // 2), Image.LANCZOS)
        im.save(out, dpi=(300, 300), optimize=True)
    return out


# ------------------------------------------------------------------ Markdown -> document
def build(md_path: Path, out_path: Path, images: bool, kind: str, line_numbers: bool = True) -> dict:
    conv = EquationConverter()
    lines = md_path.read_text(encoding="utf-8").lstrip("﻿").split("\n")
    doc = Document()
    setup(doc, line_numbers=line_numbers)
    stats = {"kind": kind, "want_line_numbers": line_numbers, "paragraphs": 0, "tables": 0, "images": 0, "headings": 0, "equations": 0, "inline_equations": 0,
             "equation_failures": [], "missing_figures": [], "references": 0, "bullets": 0}
    section = ""
    eq_no = 0
    first_table_in_section = True
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        m = FIG_RX.match(s)
        if m:
            if images:
                src = FIG / f"{m.group(1)}.png"
                if src.exists():
                    page_break(doc)
                    doc.add_picture(str(review_png(m.group(1))), width=Cm(16.0))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    doc.paragraphs[-1].paragraph_format.keep_with_next = True
                    stats["images"] += 1
                else:
                    stats["missing_figures"].append(str(src))
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            table(doc, parse_markdown_table(block), conv, stats)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level, title = len(m.group(1)), m.group(2)
            section = title
            if kind != "supplement" and title in ("参考文献", "表", "图注"):
                page_break(doc)
            if kind == "supplement" and title in ("补充表", "补充图", "补充参考文献"):
                page_break(doc)
            style = {1: "Title", 2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}[level]
            p = doc.add_paragraph(style=style)
            rich(p, title, conv, stats, size={1: 16, 2: 14, 3: 12, 4: 12}[level], bold=True, ea=EA_HEAD)
            stats["headings"] += 1
            first_table_in_section = True
            i += 1
            continue
        if s.startswith("$$"):
            block = [s]
            if not (s.endswith("$$") and len(s) > 2):
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    block.append(lines[i])
                    i += 1
                if i < len(lines):
                    block.append(lines[i])
            i += 1
            tex = " ".join(block).replace("$$", "").strip().rstrip(",，")
            eq_no += 1
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before, pf.space_after = Pt(4), Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                p._p.append(parse_xml(conv.omml(tex)))
                stats["equations"] += 1
            except Exception as exc:
                stats["equation_failures"].append({"latex": tex, "error": str(exc)})
                fonts(p.add_run(tex))
            fonts(p.add_run(f"  ({'S' if kind == 'supplement' else ''}{eq_no})"), size=SMALL_PT)
            continue
        if section in ("参考文献", "补充参考文献") and re.match(r"^\d+\.\s", s):
            para(doc, s, conv, stats, size=SMALL_PT, indent=False, hanging=21, spacing=1.2, after=2,
                 align=WD_ALIGN_PARAGRAPH.LEFT)
            stats["references"] += 1
            i += 1
            continue
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing, p.paragraph_format.space_after = 1.5, Pt(2)
            rich(p, re.sub(r"^[-*]\s+", "", s), conv, stats)
            stats["bullets"] += 1
            i += 1
            continue
        if CAPTION_RX.match(s):
            if s.startswith("表") and section in ("表", "补充表"):
                if not first_table_in_section and kind != "supplement":
                    page_break(doc)
                first_table_in_section = False
            head = s.startswith("表")
            para(doc, s, conv, stats, size=SMALL_PT, bold=head, indent=False, before=8 if head else 4, after=4,
                 keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.2)
            i += 1
            continue
        if s.startswith("注："):
            para(doc, s, conv, stats, size=TABLE_PT, indent=False, after=6, spacing=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
            i += 1
            continue
        if s.startswith("**图 "):
            para(doc, s, conv, stats, size=SMALL_PT, indent=False, before=4, after=10, spacing=1.3)
            i += 1
            continue
        if s.startswith("**短标题**") or s.startswith("**作者**") or s.startswith("¹") or s.startswith("**关键词**") or \
                (kind == "supplement" and s.startswith("**") and s.endswith("**")):
            para(doc, s, conv, stats, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, after=6)
            i += 1
            continue
        para(doc, s, conv, stats)
        i += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return verify(out_path, stats)


def verify(path: Path, stats: dict) -> dict:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        names = z.namelist()
    d = Document(str(path))
    sec = d.sections[0]
    stats.update({
        "path": str(path), "bytes": path.stat().st_size,
        "a4": abs(sec.page_width.cm - 21.0) < 0.05 and abs(sec.page_height.cm - 29.7) < 0.05,
        "line_numbers": "w:lnNumType" in xml,
        "page_field": any("PAGE" in (z_.text or "") for f in d.sections for p in f.footer.paragraphs
                          for z_ in p._p.iter(qn("w:instrText"))),
        "omml_objects": xml.count("<m:oMath>") + xml.count("<m:oMath "),
        "verified_tables": len(d.tables), "verified_images": sum(1 for n in names if n.startswith("word/media/")),
        "heading_styles": sum(1 for p in d.paragraphs if p.style.name.startswith(("Heading", "Title"))),
    })
    stats["ok"] = bool(stats["a4"] and stats["line_numbers"] == stats.get("want_line_numbers", True) and stats["page_field"]
                       and not stats["equation_failures"]
                       and not stats["missing_figures"] and stats["omml_objects"] == stats["equations"] + stats["inline_equations"])
    return stats


def run() -> dict:
    rep = {"date": DATE}
    main_md = ROOT / f"{STEM}.md"
    supp_md = ROOT / f"{SUPP_STEM}.md"
    rep["submission"] = build(main_md, ROOT / f"{STEM}.docx", images=False, kind="main")
    rep["review"] = build(main_md, ROOT / f"{REVIEW_STEM}.docx", images=True, kind="main")
    rep["supplement"] = build(supp_md, ROOT / f"{SUPP_STEM}.docx", images=True, kind="supplement")
    notes_md = ROOT / f"SARA_投稿说明_{DATE}.md"
    if notes_md.exists():
        rep["notes"] = build(notes_md, notes_md.with_suffix(".docx"), images=False, kind="notes", line_numbers=False)
    rep["all_ok"] = all(rep[k]["ok"] for k in ("submission", "review", "supplement", "notes") if k in rep)
    (ROOT / "build").mkdir(exist_ok=True)
    (ROOT / "build" / "docx_build_report9.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    return rep


if __name__ == "__main__":
    r = run()
    print(json.dumps({k: ({kk: vv for kk, vv in v.items() if kk not in ("equation_failures",)} if isinstance(v, dict) else v)
                      for k, v in r.items()}, ensure_ascii=False, indent=1))
