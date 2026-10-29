"""Checks of the English SARA package; exits non-zero when any check fails.

  1 no Chinese characters: Markdown, Word files, PDFs, figure text, public source-data CSVs;
  2 no internal / process wording (pubstyle_en.FORBIDDEN_EN); "non-inferior" only in Supplementary Note S6; no first
    person; the user's e-mail address nowhere;
  3 number tracing: every decimal, percentage and thousands-separated integer was produced by fmt9/fmt_en (or is a
    design constant); formatting: no signed zero, "nan", ASCII hyphen as minus sign, double spaces, space before
    punctuation, British spellings;
  4 the numbers of the English main text and supplement are the numbers of the Chinese second revision; the English
    source-data CSVs hold exactly the numbers of the Chinese ones;
  5 BIB/OUP references: bracketed citations numbered by first appearance, ranges only for three or more consecutive
    numbers, list numbered 1..n, entries complete (refs_en.check), no duplicates;
  6 cross references: figures, panels, tables, supplementary notes/tables/figures and equations exist; main figures and
    tables first cited in order; legend panels equal the figure panels;
  7 Word files: build audit ok, MathType objects only, one object per equation of the Markdown, bold panel letters in
    every legend, italics only where intended, round trip through MathType; PDFs: no field errors;
  8 figure palettes; 9 abbreviations defined at first use; 10 neutral headings; 11 wording boundaries (comparable,
    internal validation, same environment with a fresh cache, placeholders); abstract <= 250 words, 3-5 Key Points.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd

from . import verify9 as V9
from .build_docx_en import BUILD, REVIEW_STEM
from .manuscript_en import EN, STEM
from .mt_factory10 import DISPLAY_RX, INLINE_RX
from .numbers9 import FIG, ROOT
from .pubstyle_en import FORBIDDEN_EN_RX
from .supp_en import SUPP_STEM

CJK = re.compile(r"[\u2e80-\u9fff\u3000-\u303f\uff00-\uffef]")
FIG_EN = EN / "figures"
ZH_MAIN = ROOT / "SARA_知识覆盖不均衡条件下的疾病基因语义检索与可用性感知排序_投稿版_20260930.md"
ZH_SUPP = ROOT / "SARA_补充材料_20260930.md"
EMAIL = re.compile(r"ryuiiyreyui|[\w.+-]+@(?:gmail|qq|163|outlook|hotmail)\.com", re.I)
CONSTANTS = V9.CONSTANTS | {"19,000", "40,000"}
ABBR = ["PPI", "HPO", "GO", "RRF", "OOF", "HGNC", "ncRNA", "CTD", "MeSH", "OMIM", "HPA", "IC", "TF-IDF", "RWR", "CI",
        "SHAP", "GEO", "PMID", "IEA", "IBA"]
EVALUATIVE = re.compile(r"(?i)\b(outperform\w*|superior\w*|significant\w*|novel|improv\w*|better|best|state-of-the-art|"
                        r"breakthrough|remarkabl\w*|excellent|powerful)\b")
BRITISH = re.compile(r"(?i)\b(grey\w*|colour\w*|label{2}ed|label{2}ing|model{2}ing|centre\w*|neighbour\w*|behaviour\w*|"
                     r"favour\w*|programme|(?:optimis|normalis|characteris|summaris|visualis|randomis|minimis|maximis|"
                     r"utilis|recognis)(?:e|ed|es|ing|ation|ations)|analys(?:e|ed|ing)|towards|whilst|amongst)\b")
REF_HEAD = re.compile(r"\n## (?:Supplementary )?[Rr]eferences\n|\n## (?:补充)?参考文献\n")
TAB_HEAD = re.compile(r"\n## (?:Tables|表)\n")
# numbers as in verify9.NUM_RX, but a number may also end a sentence ("was 0.166.")
NUM_RX = re.compile(r"(?<![\w.])[+−]?\d{1,3}(?:,\d{3})+(?:\.\d+)?(?!\d|\.\d)%?|(?<![\w.])[+−]?\d+\.\d+(?!\d|\.\d)%?"
                    r"|(?<![\w.,])\d+%")


# ------------------------------------------------------------------ helpers
def body(md: str, math: bool = False, repl: str = " ") -> str:
    """text without the reference list (tables and legends kept) and, by default, with LaTeX replaced by `repl`."""
    parts = REF_HEAD.split(md)
    tab = TAB_HEAD.split(md, maxsplit=1)
    text = parts[0] + ("\n" + tab[1] if len(tab) > 1 else "")
    if not math:
        text = re.sub(r"\$\$.*?\$\$", repl, text, flags=re.S)
        text = re.sub(r"\$[^$\n]+\$", repl, text)
    return text


def refs_block(md: str) -> str:
    parts = REF_HEAD.split(md)
    return parts[1].split("\n## ", 1)[0] if len(parts) > 1 else ""


def docx_text(path: Path) -> str:
    return V9.docx_text(path)


def pdf_text(path: Path) -> str:
    sys.path.append(r"D:\信\gpu8\tools_libs10")
    import fitz
    return "\n".join(p.get_text("text") for p in fitz.open(str(path)))


def section(md: str, head: str) -> str:
    m = re.search(rf"^## {re.escape(head)}.*?$(.*?)(?=^## |\Z)", md, flags=re.M | re.S)
    return m.group(1) if m else ""


# ------------------------------------------------------------------ 1-3 text checks
def cjk(text: str, where: str) -> list[str]:
    return [f"{where}: {line[:80]}" for line in text.splitlines() if CJK.search(line)][:10]


def s6_text(text: str) -> str:
    """Supplementary Note S6 of a supplement given as Markdown or as flattened Word text."""
    s = section(text, "Supplementary Note S6")
    if s:
        return s
    m = re.search(r"^Supplementary Note S6\..*?$(.*?)^Supplementary tables$", text, flags=re.M | re.S)
    return m.group(1) if m else ""


def forbidden(md: str, where: str, allow_s6: bool = False) -> list[str]:
    out = []
    s6 = s6_text(md) if allow_s6 else ""
    for i, line in enumerate(md.splitlines(), 1):
        for m in FORBIDDEN_EN_RX.finditer(line):
            if allow_s6 and re.fullmatch(r"[Nn]on-?inferior", m.group(0)) and line.strip() and line.strip() in s6:
                continue
            out.append(f"{where}:{i}: {m.group(0)} :: {line[:90]}")
    return out


def first_person(md: str) -> list[str]:
    return [m.group(0) for m in re.finditer(r"(?i)\b(we|our|us|ours|ourselves)\b", body(md))]


def numbers(md: str, trace: dict) -> list[str]:
    bad = []
    for line in body(md).splitlines():
        s = line.strip()
        if s.startswith("#") or re.match(r"^\d+\. ", s):
            continue
        for m in V9.SCI_RX.finditer(line):
            if m.group(0) not in trace:
                bad.append(f"{m.group(0)} :: untraced scientific notation")
        line = V9.SCI_RX.sub(" ", line)
        line = re.sub(r"Sections? \d+\.\d+", " ", line)
        for m in NUM_RX.finditer(line):
            tok = m.group(0)
            if tok in trace or tok in CONSTANTS or tok.lstrip("+") in trace:
                continue
            bad.append(f"{tok} :: {line[max(0, m.start() - 30):m.end() + 20]}")
    return bad


def fmt(md: str) -> list[str]:
    text = body(md, repl="M")
    out = []
    for rx, why in ((r"[+−-]0\.0+(?![\d])", "signed zero"), (r"[+−]0%", "signed zero %"), (r"\bnan\b", "nan"),
                    (r"(?<![\w\d.–-])-\d", "ASCII hyphen as minus"), (r"(?<=\S)  +(?=\S)", "double space"),
                    (r" [,.;:)](?!\d)", "space before punctuation"), (r"\( ", "space after bracket"),
                    (r"\.\.(?!\.)", "double full stop"), (r"(?<![\w.])\d+\.\d+%?(?= ?–\d)", None)):
        if why is None:
            continue
        for m in re.finditer(rx, text):
            line = text[text.rfind("\n", 0, m.start()) + 1:m.start()]
            if why == "double space" and line.lstrip().startswith("|"):
                continue                                     # empty table cells
            out.append(f"{why}: {text[max(0, m.start() - 25):m.end() + 25]!r}")
    out += [f"British spelling: {m.group(0)}" for m in BRITISH.finditer(text)]
    return out


# ------------------------------------------------------------------ 4 consistency with the Chinese revision
def number_set(md: str) -> set[str]:
    text = body(md)
    text = V9.SCI_RX.sub(lambda m: " " + m.group(0).replace(" ", "") + " ", text)
    out = set()
    for line in text.splitlines():
        if line.strip().startswith("#"):
            continue
        line = re.sub(r"(Sections?|方法|正文方法|第) ?\d+\.\d+|\d+\.\d+ ?节", " ", line)
        out |= {m.group(0).lstrip("+") for m in NUM_RX.finditer(line)}
        out |= {m.group(0) for m in re.finditer(r"\d\.\d×10[⁻⁰¹²³⁴⁵⁶⁷⁸⁹]+", line)}
    return out


def same_numbers(en_md: str, zh_md: str) -> dict:
    en, zh = number_set(en_md), number_set(zh_md)
    allowed_en = {"19,000", "40,000"}                      # "1.9 万" / "4 万" in the Chinese text
    allowed_zh = {"1.9"}
    return {"only_en": sorted(en - zh - allowed_en), "only_zh": sorted(zh - en - allowed_zh)}


def csv_numbers() -> list[str]:
    bad = []
    zh_dir, en_dir = FIG / "source_data", FIG_EN / "source_data"
    en_files = sorted(en_dir.glob("*.csv"))
    if {p.name for p in en_files} != {p.name for p in zh_dir.glob("*.csv")}:
        bad.append("English and Chinese source-data file sets differ")
    for p in en_files:
        q = zh_dir / p.name
        if not q.exists():
            continue
        a, b = pd.read_csv(p, encoding="utf-8-sig"), pd.read_csv(q, encoding="utf-8-sig")
        if a.shape != b.shape:
            bad.append(f"{p.name}: shape {a.shape} != {b.shape}")
            continue
        for ca, cb in zip(a.columns, b.columns):
            x, y = a[ca], b[cb]
            if pd.api.types.is_numeric_dtype(y):
                if not pd.api.types.is_numeric_dtype(x) or not np.allclose(x.to_numpy(float), y.to_numpy(float),
                                                                          equal_nan=True, rtol=0, atol=0):
                    bad.append(f"{p.name}: numeric column {cb} differs")
        text = p.read_text(encoding="utf-8-sig")
        if CJK.search(text):
            bad.append(f"{p.name}: Chinese characters")
        hits = sorted({m.group(0) for m in FORBIDDEN_EN_RX.finditer(text)})
        if hits:
            bad.append(f"{p.name}: {hits[:5]}")
    return bad


# ------------------------------------------------------------------ 5 references
def citations(md: str, n_refs: int) -> list[str]:
    bad = []
    if "[@" in md:
        bad.append("unresolved [@key] citation")
    if "<sup>" in md:
        bad.append("superscript citation")
    seen = []
    for m in re.finditer(r"\[(\d+(?:(?:–|, )\d+)*)\]", body(md)):
        for part in m.group(1).split(", "):
            ends = [int(x) for x in part.split("–")]
            if len(ends) == 2 and ends[1] - ends[0] < 2:
                bad.append(f"range of fewer than three numbers: [{m.group(1)}]")
            for n in range(ends[0], ends[-1] + 1):
                if n not in seen:
                    seen.append(n)
    if seen != list(range(1, n_refs + 1)):
        bad.append(f"references not numbered by first citation: {seen[:20]}")
    entries = re.findall(r"^(\d+)\. (.+)$", refs_block(md), flags=re.M)
    if [int(n) for n, _ in entries] != list(range(1, n_refs + 1)):
        bad.append("reference list not numbered 1..n")
    texts = [e for _, e in entries]
    if len(texts) != len(set(texts)):
        bad.append("duplicate reference entries")
    for n, e in entries:
        if not e.endswith("."):
            bad.append(f"{n}: entry does not end with a full stop")
        if re.search(r"\bet al\b(?!\.)", e):
            bad.append(f"{n}: 'et al' without full stop")
    return bad


def bib_records(info: dict) -> list[str]:
    from .refs_en import CACHE, check, record
    cache = json.loads(CACHE.read_text(encoding="utf-8"))
    return check([record(k, cache) for k, _ in sorted(info["ref_order"].items(), key=lambda t: t[1])])


# ------------------------------------------------------------------ 6 cross references
FIG_REF = re.compile(r"(?:Supplementary )?Figures? (S?\d+)((?:[a-z](?:–[a-z])?)(?:, [a-z](?:–[a-z])?)*)?(?!\w|\.\d)")


def xrefs(md: str, supp: bool, n_disp: int) -> tuple[list[str], dict]:
    text = body(md)
    bad, first = [], {}
    for m in FIG_REF.finditer(text):
        n, letters = m.group(1), m.group(2) or ""
        if n not in V9.PANELS:
            bad.append(f"no figure {n}")
            continue
        first.setdefault(n, m.start())
        for part in letters.split(", "):
            if not part:
                continue
            a, b = (part.split("–") + [part])[:2]
            for c in {chr(x) for x in range(ord(a), ord(b) + 1)}:
                if c not in V9.PANELS[n]:
                    bad.append(f"figure {n} has no panel {c}")
    for m in re.finditer(r"Tables? (S?)(\d+)(?: and (S?)(\d+))?", text):
        for s, k in ((m.group(1), m.group(2)), (m.group(3), m.group(4))):
            if k is None:
                continue
            k = int(k)
            if (s and not 1 <= k <= 10) or (not s and not 1 <= k <= 3):
                bad.append(f"no table {s}{k}")
            first.setdefault(f"T{s}{k}", m.start())
    for m in re.finditer(r"Notes? S(\d+)", text):
        if not 1 <= int(m.group(1)) <= 6:
            bad.append(f"no note S{m.group(1)}")
    for m in re.finditer(r"Eq\. \((S?)(\d+)\)", text):
        if (m.group(1) == "S") != supp or not 1 <= int(m.group(2)) <= n_disp:
            bad.append(f"Eq. ({m.group(1)}{m.group(2)}) has no equation")
    if "@eq:" in md or "{#eq:" in md:
        bad.append("unresolved equation label")
    if not supp:
        nums = [int(k) for _, k in sorted((v, k) for k, v in first.items() if k.isdigit())]
        if nums != sorted(nums):
            bad.append(f"main figures cited out of order: {nums}")
        torder = [int(k[1:]) for _, k in sorted((v, k) for k, v in first.items() if re.fullmatch(r"T\d+", k))]
        if torder != sorted(torder):
            bad.append(f"main tables cited out of order: {torder}")
    return bad, first


def legend_panels(md: str) -> list[str]:
    bad = []
    for m in re.finditer(r"^\*\*Figure (S?\d+)\. [^*]+\*\* (.+)$", md, flags=re.M):
        n, leg = m.group(1), m.group(2)
        got = "".join(sorted(set(re.findall(r"\*\*([a-z])\*\*", leg))))
        want = V9.PANELS[n] if n not in ("10",) else V9.PANELS[n]
        if got != want:
            bad.append(f"legend of figure {n}: panels {got} != {want}")
    return bad


def equation_counts(md: str) -> dict:
    lines = md.split("\n")
    return {"display": sum(1 for s in lines if DISPLAY_RX.match(s.strip())),
            "inline": sum(len(INLINE_RX.findall(s)) for s in lines if not DISPLAY_RX.match(s.strip()))}


# ------------------------------------------------------------------ 7 Word files
def docx_runs(path: Path) -> list[str]:
    """legend panel letters must be bold runs; italic runs may only be the P of P values."""
    from docx import Document
    bad = []
    d = Document(str(path))
    for p in d.paragraphs:
        m = re.match(r"Figure (S?\d+)\. ", p.text)
        if m and len(p.text) > 200:
            letters = "".join(sorted({r.text.strip() for r in p.runs if r.bold and re.fullmatch(r"[a-z]", r.text.strip())}))
            if letters != V9.PANELS[m.group(1)]:
                bad.append(f"{path.name}: Figure {m.group(1)} bold panels {letters}")
        for r in p.runs:
            if r.italic and r.text.strip() and r.text.strip() != "P":
                bad.append(f"{path.name}: unexpected italic run {r.text[:40]!r}")
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        if r.italic and r.text.strip() and r.text.strip() != "P":
                            bad.append(f"{path.name}: unexpected italic run in table {r.text[:40]!r}")
    return sorted(set(bad))[:20]


def docx_report(cnt_m: dict, cnt_s: dict) -> list[str]:
    bad = []
    p = BUILD / "docx_build_report_en.json"
    if not p.exists():
        return ["docx build report missing"]
    b = json.loads(p.read_text(encoding="utf-8"))
    if not b.get("all_ok"):
        bad.append("docx build report not ok")
    for key, cnt in (("submission", cnt_m), ("review", cnt_m), ("supplement", cnt_s)):
        d = b.get(key, {})
        if d.get("omml_objects") != 0:
            bad.append(f"{key}: OMML objects present")
        if (d.get("display_equations"), d.get("inline_equations")) != (cnt["display"], cnt["inline"]):
            bad.append(f"{key}: equations {d.get('display_equations')}/{d.get('inline_equations')} != Markdown "
                       f"{cnt['display']}/{cnt['inline']}")
        if d.get("ole_objects") != cnt["display"] + cnt["inline"]:
            bad.append(f"{key}: {d.get('ole_objects')} MathType objects for {cnt['display'] + cnt['inline']} equations")
        if key != "submission" and d.get("missing_figures"):
            bad.append(f"{key}: missing figures")
    exp = {"submission": 0, "review": 10, "supplement": 8}
    for key, n in exp.items():
        if b.get(key, {}).get("images") != n:
            bad.append(f"{key}: {b.get(key, {}).get('images')} figures, expected {n}")
    return bad


def roundtrip() -> list[str]:
    p = BUILD / "roundtrip_en.json"
    if not p.exists():
        return ["MathType round trip not run"]
    r = json.loads(p.read_text(encoding="utf-8"))
    return [] if r.get("ok") else [f"MathType round trip failed: {r}"]


def pdf_checks() -> tuple[list[str], dict]:
    bad, pages = [], {}
    rep = BUILD / "word_pdf_report_en.json"
    if not rep.exists():
        return ["Word PDF export report missing"], pages
    r = json.loads(rep.read_text(encoding="utf-8"))
    for key, d in r.items():
        p = Path(d["pdf"])
        if not p.exists():
            bad.append(f"{key}: PDF missing")
            continue
        t = pdf_text(p)
        pages[key] = d.get("pages")
        bad += cjk(t, f"pdf:{key}")
        if re.search(r"Error! |Reference source not found|Bookmark not defined|MERGEFORMAT", t):
            bad.append(f"{key}: field error in PDF")
        if d.get("blank_pages"):
            bad.append(f"{key}: blank pages {d['blank_pages']}")
        if EMAIL.search(t):
            bad.append(f"{key}: e-mail address in PDF")
    return bad, pages


# ------------------------------------------------------------------ 9-11
def abbreviations(md: str) -> list[str]:
    text = md.split("## 1 Introduction", 1)[1].split("## References", 1)[0]
    text = re.sub(r"\$[^$\n]+\$", " ", re.sub(r"\$\$.*?\$\$", " ", text, flags=re.S))
    bad = []
    for a in ABBR:
        m = re.search(rf"(?<![A-Za-z]){re.escape(a)}s?(?![A-Za-z])", text)
        if not m:
            bad.append(f"{a}: not used in the main text")
            continue
        after = text[m.end():m.end() + 1]
        before = text[max(0, m.start() - 160):m.start()]
        if not (after == ")" and "(" in before and ")" not in before[before.rfind("("):]):
            bad.append(f"{a}: first use not defined :: {text[max(0, m.start() - 50):m.end() + 10]}")
    return bad


def titles(md: str) -> list[str]:
    bad = []
    for line in md.split("\n"):
        s = line.strip()
        m = re.match(r"^#{2,4} (?:\d+(?:\.\d+)* |Supplementary Note S\d+\. )?(.+)$", s)
        if m:
            if EVALUATIVE.search(m.group(1)):
                bad.append(f"evaluative heading: {m.group(1)}")
            if len(m.group(1)) > 110:
                bad.append(f"heading longer than 110 characters: {m.group(1)}")
        m = re.match(r"^(?:\*\*)?(Figure|Table) S?\d+\. ([^*]+?)(?:\.\*\*|$)", s)
        if m and EVALUATIVE.search(m.group(2)):
            bad.append(f"evaluative {m.group(1)} title: {m.group(2)}")
    return bad


def boundaries(main: str, supp: str) -> list[str]:
    bad = []
    need_main = {
        "direct fusion comparable": r"Compared with multimodal direct fusion, which uses[\s\S]{0,300}?the two methods "
                                    r"performed comparably",
        "abstract comparable": r"SARA performed comparably to a multimodal direct-fusion model",
        "temporal comparable": r"SARA and retrieval-channel LightGBM performed comparably on the new-association task",
        "internal validation": r"the cross-validation folds and the external label set were used during the iterative "
                               r"development of the method, so the results constitute retrospective internal-validation",
        "same environment, fresh cache": r"In the same computing environment, the complete pipeline[^.]*?with a fresh cache",
        "placeholder authors": r"\[Author name\]", "placeholder funding": r"\[Funding bodies and grant numbers\]",
        "placeholder conflicts": r"\[The authors declare no conflict of interest",
        "placeholder contributions": r"\[Author A: conceptualization",
        "placeholder AI disclosure": r"\[If artificial intelligence tools were used",
        "placeholder biographies": r"\[Author name\] is a \[position\]",
        "placeholder data links": r"\[code repository link\] and \[data archive DOI\]"}
    for k, rx in need_main.items():
        if not re.search(rx, main, flags=re.S):
            bad.append(f"missing: {k}")
    if re.search(r"[Nn]on-?inferior", main):
        bad.append("non-inferiority wording in the main text")
    s6 = s6_text(supp)
    if not re.search(r"non-inferiority margin of −0\.005", s6):
        bad.append("non-inferiority note missing from Supplementary Note S6")
    if re.search(r"[Nn]on-?inferior", supp.replace(s6, "")):
        bad.append("non-inferiority wording outside Supplementary Note S6")
    if re.search(r"(?i)multimodal direct fusion[^.]{0,80}\b(outperform|superior|exceed)", main):
        bad.append("SARA claimed better than multimodal direct fusion")
    return bad


# ------------------------------------------------------------------ run
def run() -> dict:
    main = (EN / f"{STEM}.md").read_text(encoding="utf-8")
    supp = (EN / f"{SUPP_STEM}.md").read_text(encoding="utf-8")
    info_m = json.loads((BUILD / "manuscript_en_info.json").read_text(encoding="utf-8"))
    info_s = json.loads((BUILD / "supp_en_info.json").read_text(encoding="utf-8"))
    docs = {"submission": EN / f"{STEM}.docx", "review": EN / f"{REVIEW_STEM}.docx", "supplement": EN / f"{SUPP_STEM}.docx"}
    rep: dict = {}
    rep["1_cjk_markdown"] = cjk(main, "main") + cjk(supp, "supp")
    rep["1_cjk_docx"] = sum((cjk(docx_text(p), k) if p.exists() else [f"{k}: missing"] for k, p in docs.items()), [])
    scan = json.loads((FIG_EN / "text_scan.json").read_text(encoding="utf-8"))
    rep["1_figures"] = [f"{k}: cjk {v.get('cjk')} forbidden {v.get('forbidden')}" for k, v in scan.items()
                        if isinstance(v, dict) and (v.get("cjk") or v.get("forbidden"))]
    if len([k for k, v in scan.items() if isinstance(v, dict)]) != 18:
        rep["1_figures"].append("figure text scan does not cover 18 figures")
    rep["2_forbidden"] = (forbidden(main, "main") + forbidden(supp, "supp", allow_s6=True) +
                          sum((forbidden(docx_text(p), k, allow_s6=(k == "supplement")) for k, p in docs.items()
                               if p.exists()), []))
    rep["2_first_person"] = first_person(main) + first_person(supp)
    rep["2_email"] = [k for k, t in (("main", main), ("supp", supp)) if EMAIL.search(t)] + \
        [k for k, p in docs.items() if p.exists() and EMAIL.search(docx_text(p))]
    rep["3_untraced_main"] = numbers(main, info_m["trace"])
    rep["3_untraced_supp"] = numbers(supp, info_s["trace"])
    rep["3_format"] = fmt(main) + fmt(supp)
    sm, ss = same_numbers(main, ZH_MAIN.read_text(encoding="utf-8")), same_numbers(supp, ZH_SUPP.read_text(encoding="utf-8"))
    rep["4_numbers_vs_chinese"] = [f"main {k}: {v}" for k, v in sm.items() if v] + [f"supp {k}: {v}" for k, v in ss.items() if v]
    rep["4_csv_numbers"] = csv_numbers()
    rep["5_citations_main"] = citations(main, info_m["n_refs"])
    rep["5_citations_supp"] = citations(supp, info_s["n_refs"])
    rep["5_bib_main"] = bib_records(info_m)
    rep["5_bib_supp"] = bib_records(info_s)
    cnt_m, cnt_s = equation_counts(main), equation_counts(supp)
    bm, first = xrefs(main, False, cnt_m["display"])
    bs, _ = xrefs(supp, True, cnt_s["display"])
    rep["6_xrefs"] = bm + bs
    rep["6_legend_panels"] = legend_panels(main) + legend_panels(supp)
    rep["7_docx"] = docx_report(cnt_m, cnt_s)
    rep["7_docx_runs"] = sum((docx_runs(p) for k, p in docs.items() if k != "submission" and p.exists()), [])
    rep["7_roundtrip"] = roundtrip()
    rep["7_pdf"], pages = pdf_checks()
    man = json.loads((FIG_EN / "figures_manifest.json").read_text(encoding="utf-8"))
    pal = man.get("_palette_checks", {})
    rep["8_palette"] = [k for k, v in pal.items() if isinstance(v, dict) and not v.get("pass", True)] if pal else ["no palette checks"]
    rep["9_abbreviations"] = abbreviations(main)
    rep["10_titles"] = titles(main) + titles(supp)
    rep["11_boundaries"] = boundaries(main, supp)
    cover = BUILD / "SARA_cover_letter_20260930.md"
    if cover.exists():
        ct = cover.read_text(encoding="utf-8")
        trace_c = json.loads((BUILD / "cover_en_info.json").read_text(encoding="utf-8"))["trace"]
        rep["12_cover_letter"] = (cjk(ct, "cover") + forbidden(ct, "cover") + numbers(ct, trace_c) +
                                  (["e-mail address in the cover letter"] if EMAIL.search(ct) else []) +
                                  ([] if (EN / "SARA_cover_letter_20260930.docx").exists() else ["cover letter docx missing"]))
    else:
        rep["12_cover_letter"] = ["cover letter missing"]
    gl = EN / "SARA_术语对照_20260930.xlsx"
    if gl.exists():
        from openpyxl import load_workbook
        ws = load_workbook(str(gl)).active
        eng = [str(r[2]) for r in ws.iter_rows(min_row=2, values_only=True)]
        rep["12_glossary"] = ([f"Chinese in English column: {e}" for e in eng if CJK.search(e)] +
                              [f"forbidden in glossary: {e}" for e in eng if FORBIDDEN_EN_RX.search(e)])
        rep["12_glossary"] += [] if len(eng) >= 150 else [f"glossary has only {len(eng)} terms"]
    else:
        rep["12_glossary"] = ["glossary missing"]
    abstract = re.search(r"## Abstract\n\n(.+?)\n\n", main, flags=re.S).group(1)
    words = len(abstract.split())
    kp = section(main, "Key Points").count("\n- ")
    rep["11_lengths"] = ([f"abstract has {words} words"] if words > 250 else []) + \
        ([f"{kp} Key Points"] if not 3 <= kp <= 5 else [])
    rep["_info"] = {"abstract_words": words, "key_points": kp, "main_words": len(main.split("## References")[0].split()),
                    "refs_main": info_m["n_refs"], "refs_supp": info_s["n_refs"], "equations_main": cnt_m,
                    "equations_supp": cnt_s, "pdf_pages": pages, "numbers_main": len(number_set(main)),
                    "numbers_supp": len(number_set(supp)),
                    "first_mentions": {k: v for k, v in sorted(first.items(), key=lambda t: t[1])}}
    fails = {k: v for k, v in rep.items() if not k.startswith("_") and v}
    rep["ok"] = not fails
    BUILD.mkdir(parents=True, exist_ok=True)
    (BUILD / "verify_en_report.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    return rep


if __name__ == "__main__":
    r = run()
    show = {k: (v[:12] if isinstance(v, list) else v) for k, v in r.items() if k != "_info"}
    print(json.dumps(show, ensure_ascii=False, indent=1))
    print(json.dumps({k: v for k, v in r["_info"].items() if k != "first_mentions"}, ensure_ascii=False))
    sys.exit(0 if r["ok"] else 1)
