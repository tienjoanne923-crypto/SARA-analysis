"""Chinese-English terminology table of the SARA package (Word and Excel).

Built from the term sources of the English package, so the table lists exactly the wording used in the English text,
tables and figures: terms_en (methods, design, audit terms, features, experimental arms, inputs, case diseases),
figtext_en (MeSH categories, GO evidence classes) and the paired headings, figure titles and table titles of the Chinese
second revision and the English version.
"""
from __future__ import annotations

import json
import re

from . import supp9 as S
from . import terms_en as T
from .figtext_en import EVIDENCE, MESH
from .manuscript9 import CASE_ZH
from .manuscript10 import FIGS10
from .manuscript_en import DATE, EN, FIG_TITLES
from .numbers9 import ROOT
from .pubdata8 import FEAT_ZH
from .publabel9 import relabel
from .supp_en import SFIGS

STEM = f"SARA_术语对照_{DATE}"
BUILD = ROOT / "build" / "en"
ZH_MAIN = ROOT / "SARA_知识覆盖不均衡条件下的疾病基因语义检索与可用性感知排序_投稿版_20260930.md"
ZH_SUPP = ROOT / "SARA_补充材料_20260930.md"


def _headings(md: str) -> list[str]:
    return [m.group(2) for m in re.finditer(r"^(#{2,3}) (.+)$", md, flags=re.M)]


def _strip_num(h: str) -> str:
    return re.sub(r"^(\d+(?:\.\d+)* |补充说明 S\d+　|Supplementary Note S\d+\. )", "", h)


def groups() -> list[tuple[str, list[tuple[str, str]]]]:
    out = [(g, list(pairs)) for g, pairs in T.GLOSSARY]
    out.append(("排序特征族（表 S2）", [(S.FAM_PUB[k], v) for k, v in T.FAMILY.items()]))
    out.append(("排序特征（表 S2）", [(relabel(FEAT_ZH[f]), T.cap(T.FEATURE[f])) for f in T.FEATURE]))
    out.append(("证据删除条件（表 S7，图 8）", [(zh, en) for (_, zh, _), (_, en, _) in zip(S.E4_PUB, T.E4)]))
    out.append(("序列补位方式（表 S8）", [(zh, en) for (_, zh), (_, en) in zip(S.E5_PUB, T.E5)]))
    out.append(("支持任务方法（表 S10）", [(S.SUP_PUB[k], T.SUPPORT[k]) for k in S.SUP_PUB]))
    out.append(("全背景语义排名（表 S3）", [(S.SEM_PUB[k], T.SEMANTIC[k]) for k in S.SEM_PUB]))
    out.append(("时间外推的输入（表 S9）", [(S.HIST_PUB[k], T.HIST[k]) for k in S.HIST_PUB]))
    out.append(("MeSH 疾病大类", [(zh, T.cap(v[1])) for zh, v in MESH.items()]))
    out.append(("案例疾病（图 10）", [(CASE_ZH[k], T.cap(v)) for k, v in T.CASE.items()]))
    out.append(("GO 证据类别（图 8a）", [(zh, en) for zh, en in EVIDENCE.items()]))
    zm, em = ZH_MAIN.read_text(encoding="utf-8"), (EN / f"SARA_manuscript_{DATE}.md").read_text(encoding="utf-8")
    zs, es = ZH_SUPP.read_text(encoding="utf-8"), (EN / f"SARA_supplementary_{DATE}.md").read_text(encoding="utf-8")
    heads = []
    for z, e in ((zm, em), (zs, es)):
        hz, he = _headings(z), _headings(e)
        assert len(hz) == len(he), (len(hz), len(he))
        heads += [(_strip_num(a), _strip_num(b)) for a, b in zip(hz, he)]
    out.append(("章节标题", heads))
    figs = [(t, FIG_TITLES[n]) for n, _, t in FIGS10]
    zt = dict(re.findall(r"\*\*图 (S\d+)　([^。*]+)。\*\*", zs))
    figs += [(zt[n], t) for n, _, t, _ in SFIGS]
    out.append(("图题", figs))
    tz = dict(re.findall(r"^表 (S?\d+)　(.+)$", zm + "\n" + zs, flags=re.M))
    te = dict(re.findall(r"^Table (S?\d+)\. (.+)$", em + "\n" + es, flags=re.M))
    assert set(tz) == set(te), (sorted(tz), sorted(te))
    order = sorted(tz, key=lambda k: (k.startswith("S"), int(k.lstrip("S"))))
    out.append(("表题", [(tz[k], te[k]) for k in order]))
    return out


def _plain(s: str) -> str:
    return re.sub(r"\$([^$]+)\$", lambda m: m.group(1).replace("\\", ""), s)


def write() -> dict:
    gs = groups()
    parts = ["# SARA 中英术语对照表",
             "本表由英文版的术语来源自动导出，与英文正文、表格和图中的用词一致；数学符号在 Word 中为 MathType 公式。"]
    rows_x = []
    for g, pairs in gs:
        seen, uniq = set(), []
        for zh, en in pairs:
            if (zh, en) not in seen:
                seen.add((zh, en))
                uniq.append((zh, en))
        parts.append(f"## {g}")
        parts.append("| 中文 | English |\n|---|---|\n" + "\n".join(f"| {zh} | {en} |" for zh, en in uniq))
        rows_x += [(g, _plain(zh), _plain(en)) for zh, en in uniq]
    md = "\n\n".join(parts) + "\n"
    BUILD.mkdir(parents=True, exist_ok=True)
    p = BUILD / f"{STEM}.md"
    p.write_text(md, encoding="utf-8")
    from .build_docx10 import Builder, load_assets
    b = Builder("notes", load_assets([p]), line_numbers=False)
    st = b.build(p, EN / f"{STEM}.docx", images=False)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "术语对照"
    ws.append(["类别", "中文", "English"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F1")
    for r in rows_x:
        ws.append(list(r))
    for col, w in zip("ABC", (26, 48, 70)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{ws.max_row}"
    wb.properties.title = "SARA terminology"
    wb.properties.creator = ""
    wb.save(str(EN / f"{STEM}.xlsx"))
    rep = {"docx": st["path"], "ok": st["ok"], "groups": len(gs), "terms": len(rows_x),
           "missing_assets": st["missing_assets"]}
    (BUILD / "glossary_en.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    return rep


if __name__ == "__main__":
    print(json.dumps(write(), ensure_ascii=False, indent=1))
