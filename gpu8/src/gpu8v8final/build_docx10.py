"""Word documents of the SARA submission, second revision.

Built from the Markdown written by ``manuscript10`` / ``supp10`` / ``notes10``:
  * every equation is a MathType object: the Equation.DSMT4 OLE storage and WMF preview made by MathType itself from
    the LaTeX source (``mt_factory10``); objects set in 9 pt tables or 10.5 pt captions and legends carry MathType
    preferences scaled to that size, so they keep their size when edited in MathType;
  * display equations use MathType's own layout (paragraph style MTDisplayEquation: centre tab for the equation, right
    tab for the number); numbers are ``SEQ MTEqn`` fields inside ``ZEqnNum`` bookmarks and every "式（n）" in the text
    is a ``REF`` field to that bookmark, so MathType's "Update Equation Numbers" and Word's field update keep them right;
  * GB/T 7714 citations are superscript and every number is hyperlinked to the bookmark of its reference entry; DOIs
    and URLs in the reference list are external hyperlinks;
  * the MTWinEqns custom property marks the file as a MathType document.
No OMML is written.  Every saved file is reopened and audited; the report goes to ``build/docx_build_report10.json``.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips
from lxml import etree

from . import mt_factory10 as MT
from .build_docx9 import (BODY_PT, EA_BODY, EA_HEAD, FIG_RX, SMALL_PT, TABLE_PT, _fonts, parse_markdown_table,
                          review_png, set_cell_border, setup)
from .manuscript10 import DATE, STEM, TITLE
from .numbers9 import FIG, ROOT
from .supp10 import SUPP_STEM

REVIEW_STEM = STEM.replace("投稿版", "审阅版")
NOTES_STEM = f"SARA_投稿说明_{DATE}"
TEXT_W = 9026                                      # twips between the margins of an A4 page (21 cm − 2 × 2.54 cm)
ESC = ""
NS_DECL = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
SHAPETYPE = (
    '<v:shapetype id="_x0000_t75" coordsize="21600,21600" o:spt="75" o:preferrelative="t" '
    'path="m@4@5l@4@11@9@11@9@5xe" filled="f" stroked="f"><v:stroke joinstyle="miter"/><v:formulas>'
    '<v:f eqn="if lineDrawn pixelLineWidth 0"/><v:f eqn="sum @0 1 0"/><v:f eqn="sum 0 0 @1"/><v:f eqn="prod @2 1 2"/>'
    '<v:f eqn="prod @3 21600 pixelWidth"/><v:f eqn="prod @3 21600 pixelHeight"/><v:f eqn="sum @0 0 1"/>'
    '<v:f eqn="prod @6 1 2"/><v:f eqn="prod @7 21600 pixelWidth"/><v:f eqn="sum @8 21600 0"/>'
    '<v:f eqn="prod @7 21600 pixelHeight"/><v:f eqn="sum @10 21600 0"/></v:formulas>'
    '<v:path o:extrusionok="f" gradientshapeok="t" o:connecttype="rect"/><o:lock v:ext="edit" aspectratio="t"/>'
    '</v:shapetype>')
CUSTOM_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
              'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
              '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="MTWinEqns"><vt:bool>true</vt:bool>'
              '</property></Properties>')
TOKEN = re.compile(r"(\$[^$]+\$|<sup>\[[^\]]+\]</sup>|\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+?`|式（S?\d+）)")
PVAL = re.compile(r"(?<![A-Za-z])(P)(?= [=<])")
CAPTION_RX = re.compile(r"^(表 S?\d+　|（[a-z]）)")
REF_RX = re.compile(r"^\[(\d+)\]\s+(.*)$")
LINK_RX = re.compile(r"(DOI:(10\.\S+)|(https?://\S+))\.$")


def _pt(x: float) -> str:
    return f"{x:.2f}".rstrip("0").rstrip(".")


def _cw(c: str) -> float:
    """advance width of one character of 9 pt Times New Roman / SimSun in half-em units."""
    if ord(c) >= 0x2E80:
        return 2.0
    if c == "@":
        return 1.9
    if c == "%":
        return 1.7
    if c.isupper():
        return 1.45
    if c == " ":
        return 0.5
    if c in ".,:;'|()[]!/":
        return 0.6
    return 1.0


def _plain(s: str) -> str:
    s = re.sub(r"\$([^$]+)\$", lambda m: "x" * max(1, len(re.sub(r"\\[a-zA-Z]+|[{}^_\s]", "", m.group(1)))), s or "")
    return re.sub(r"<sup>.*?</sup>", "xx", s).replace("**", "").replace("`", "")


def _units(s: str) -> float:
    return sum(_cw(c) for c in _plain(s))


def _token(s: str) -> float:
    """widest piece that cannot be broken (a Latin word or number; CJK text breaks between characters)."""
    return max((sum(_cw(c) for c in t) for t in re.findall(r"[^\s\u2e80-\uffff（）]+", _plain(s))), default=0) + 1


UNIT, PAD = 90, 260                                          # twips per half-em at 9 pt; cell margins and slack


def column_widths(rows, ncol: int, total: int = TEXT_W) -> list[int]:
    """twips per column: every column keeps its longest unbreakable token; if the table fits, nothing wraps; otherwise
    headers may wrap to two lines, and if that is not enough the widest columns wrap to a common cap."""
    cells = [[r[j] if j < len(r) else "" for r in rows] for j in range(ncol)]
    body = [max((_units(c) for c in col[1:]), default=0) for col in cells]
    head = [_units(col[0]) * 1.08 for col in cells]
    floor = [max(_token(c) * (1.08 if i == 0 else 1) for i, c in enumerate(col)) * UNIT + PAD for col in cells]
    full = [max(b, h, 4) * UNIT + PAD for b, h in zip(body, head)]
    if sum(full) <= total:
        return [int(w * total / sum(full)) for w in full]
    nat = [max(max(b, h / 2, 4) * UNIT + PAD, f) for b, h, f in zip(body, head, floor)]
    if sum(nat) <= total:
        return [int(w * total / sum(nat)) for w in nat]
    lo, hi = 0.0, max(nat)
    for _ in range(60):
        c = (lo + hi) / 2
        lo, hi = (c, hi) if sum(max(min(w, c), f) for w, f in zip(nat, floor)) <= total else (lo, c)
    return [int(max(min(w, lo), f)) for w, f in zip(nat, floor)]


NUMERIC_CELL = re.compile(r"[+−\-]?[\d.,]+%?(（[^）]*）)?"                   # a value, optionally with its interval
                          r"|[+−\-]?[\d.,]+%?( ?[/±] ?[+−\-]?[\d.,]+%?)+"     # a / b (/ c), a ± b
                          r"|\d{4}-\d{2}(-\d{2})?|—|参照|是|否|\S*\d\S*（\*{0,2}.*）")


def numeric_column(col: list[str]) -> bool:
    cells = [c.strip() for c in col[1:] if c.strip()]
    return bool(cells) and sum(bool(NUMERIC_CELL.fullmatch(c.replace("**", ""))) for c in cells) >= 0.6 * len(cells)


class Builder:
    def __init__(self, kind: str, assets: dict, line_numbers: bool = True):
        self.kind, self.assets = kind, assets
        self.doc = Document()
        setup(self.doc, line_numbers=line_numbers)
        self._styles()
        cp = self.doc.core_properties
        cp.title, cp.author, cp.last_modified_by, cp.comments, cp.keywords = TITLE, "", "", "", ""
        self.eq_prefix = "S" if kind == "supplement" else ""
        self.bm_base = 900000 if kind == "supplement" else 800000
        self.n_obj = self.bm_id = self.eq_no = 0
        self.requests = []                                  # (tex, mode, text size) of every object placed
        self.stats = {"kind": kind, "want_line_numbers": line_numbers, "paragraphs": 0, "tables": 0, "images": 0,
                      "headings": 0, "display_equations": 0, "inline_equations": 0, "objects": 0, "citations": 0,
                      "cited_numbers": [], "eq_refs": [], "references": 0, "bullets": 0, "missing_figures": [],
                      "missing_assets": [], "sizes": {}}

    # -------------------------------------------------------------- styles
    def _styles(self):
        st = self.doc.styles.add_style("MTDisplayEquation", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = self.doc.styles["Normal"]
        pf = st.paragraph_format
        pf.tab_stops.add_tab_stop(Twips(TEXT_W // 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Twips(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
        pf.first_line_indent = Pt(0)
        pf.space_before, pf.space_after, pf.line_spacing = Pt(6), Pt(6), 1.0
        pf.keep_together = True
        # Use Word's semantic Caption style for figure/table captions so that
        # captions remain consistently styled and can be targeted by journal
        # templates or accessibility tooling after export.
        cap = self.doc.styles["Caption"]
        cap.font.name = "Times New Roman"
        cap._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        cap.font.size = Pt(SMALL_PT)
        cap.font.bold = False
        cap.paragraph_format.space_before = Pt(4)
        cap.paragraph_format.space_after = Pt(6)
        cap.paragraph_format.line_spacing = 1.3
        cap.paragraph_format.keep_with_next = True

    # -------------------------------------------------------------- runs
    def run(self, p, text, size=BODY_PT, bold=False, italic=False, ea=EA_BODY, mono=False, sup=False):
        r = p.add_run(text.replace(ESC, "*"))
        _fonts(r, ea=ea, size=size, bold=bold, italic=italic, mono=mono)
        if sup:
            r.font.superscript = True
        return r

    def plain(self, p, text, size, bold, italic, ea):
        for k, seg in enumerate(PVAL.split(text)):
            if seg:
                self.run(p, seg, size, bold, italic or k % 2 == 1, ea)

    def rich(self, p, text, size=BODY_PT, bold=False, italic=False, ea=EA_BODY):
        for tok in TOKEN.split((text or "").replace("\\*", ESC)):
            if not tok:
                continue
            if tok.startswith("$") and tok.endswith("$") and len(tok) > 2:
                self.inline_math(p, tok[1:-1], size)
            elif tok.startswith("<sup>["):
                self.citation(p, tok[len("<sup>["):-len("]</sup>")], size, bold, ea)
            elif tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                self.rich(p, tok[2:-2], size, True, italic, ea)
            elif tok.startswith("`") and tok.endswith("`"):
                self.run(p, tok[1:-1], size - 1, bold, italic, ea, mono=True)
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                self.rich(p, tok[1:-1], size, bold, True, ea)
            elif tok.startswith("式（"):
                self.eq_ref(p, tok, size, bold, ea)
            else:
                self.plain(p, tok, size, bold, italic, ea)
        return p

    def field(self, p, instr, result, size=BODY_PT, bold=False, ea=EA_BODY):
        def fld(kind):
            r = self.run(p, "", size, bold, ea=ea)
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
        fld("begin")
        r = self.run(p, "", size, bold, ea=ea)
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {instr} "
        r._r.append(it)
        fld("separate")
        self.run(p, result, size, bold, ea=ea)
        fld("end")

    def bookmark(self, p, name, start: bool):
        if start:
            self.bm_id += 1
            self._open_bm = self.bm_id
        el = OxmlElement("w:bookmarkStart" if start else "w:bookmarkEnd")
        el.set(qn("w:id"), str(self._open_bm))
        if start:
            el.set(qn("w:name"), name)
        p._p.append(el)

    def link(self, p, text, size, bold=False, ea=EA_BODY, anchor=None, url=None, sup=False):
        h = OxmlElement("w:hyperlink")
        if anchor:
            h.set(qn("w:anchor"), anchor)
        else:
            h.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
        h.set(qn("w:history"), "1")
        r = self.run(p, text, size, bold, ea=ea, sup=sup)
        r._r.addprevious(h)
        h.append(r._r)

    # -------------------------------------------------------------- citations and equation references
    def citation(self, p, inner, size, bold, ea):
        self.run(p, "[", size, bold, ea=ea, sup=True)
        for k, part in enumerate(inner.split(",")):
            if k:
                self.run(p, ",", size, bold, ea=ea, sup=True)
            for m, n in enumerate(part.split("-")):
                if m:
                    self.run(p, "-", size, bold, ea=ea, sup=True)
                self.link(p, n, size, bold, ea, anchor=f"ref_{int(n)}", sup=True)
                self.stats["cited_numbers"].append(int(n))
        self.run(p, "]", size, bold, ea=ea, sup=True)
        self.stats["citations"] += 1

    def eq_bm(self, n: int) -> str:
        return f"ZEqnNum{self.bm_base + n:06d}"

    def eq_ref(self, p, tok, size, bold, ea):
        m = re.fullmatch(r"式（(S?)(\d+)）", tok)
        n = int(m.group(2))
        self.run(p, "式（" + m.group(1), size, bold, ea=ea)
        self.field(p, f"REF {self.eq_bm(n)} \\h \\* MERGEFORMAT", str(n), size, bold, ea)
        self.run(p, "）", size, bold, ea=ea)
        self.stats["eq_refs"].append(n)

    # -------------------------------------------------------------- MathType objects
    def _object(self, tex: str, mode: str, size: float):
        info = self.assets.get((tex.strip(), mode))
        if info is None:
            self.stats["missing_assets"].append((mode, tex))
            return None
        self.requests.append((tex.strip(), mode, size))
        self.stats["sizes"][str(size)] = self.stats["sizes"].get(str(size), 0) + 1
        return self.object_from_asset(MT.sized(info, size / 12.0))

    def object_from_asset(self, a: dict):
        part = self.doc.part
        pkg = part.package
        self.n_obj += 1
        n = self.n_obj
        img = Part(PackURI(f"/word/media/mteq{n}{a['image_ext']}"), CT.X_WMF, Path(a["image"]).read_bytes(), pkg)
        ole = Part(PackURI(f"/word/embeddings/oleObject{n}.bin"), CT.OFC_OLE_OBJECT, Path(a["bin"]).read_bytes(), pkg)
        rid_img, rid_ole = part.relate_to(img, RT.IMAGE), part.relate_to(ole, RT.OLE_OBJECT)
        sid = 1024 + n
        xml = (f'<w:r {NS_DECL}><w:rPr><w:position w:val="{a["position"]}"/></w:rPr>'
               f'<w:object w:dxaOrig="{a["dxa"]}" w:dyaOrig="{a["dya"]}">{SHAPETYPE if n == 1 else ""}'
               f'<v:shape id="_x0000_i{sid}" type="#_x0000_t75" style="width:{_pt(a["width_pt"])}pt;'
               f'height:{_pt(a["height_pt"])}pt" o:ole=""><v:imagedata r:id="{rid_img}" o:title=""/></v:shape>'
               f'<o:OLEObject Type="Embed" ProgID="Equation.DSMT4" ShapeID="_x0000_i{sid}" DrawAspect="Content" '
               f'ObjectID="_{1852300000 + n}" r:id="{rid_ole}"/></w:object></w:r>')
        self.stats["objects"] += 1
        return parse_xml(xml)

    def inline_math(self, p, tex, size):
        el = self._object(tex, "inline", size)
        if el is None:
            self.run(p, f"${tex}$", size)
            return
        p._p.append(el)
        self.stats["inline_equations"] += 1

    def display(self, tex):
        self.eq_no += 1
        p = self.doc.add_paragraph(style="MTDisplayEquation")
        self.run(p, "").add_tab()
        el = self._object(tex, "display", BODY_PT)
        if el is None:
            self.run(p, f"$${tex}$$")
        else:
            p._p.append(el)
        self.run(p, "").add_tab()
        self.run(p, "(" + self.eq_prefix)
        self.bookmark(p, self.eq_bm(self.eq_no), True)
        self.field(p, "SEQ MTEqn \\* ARABIC \\* MERGEFORMAT", str(self.eq_no))
        self.bookmark(p, "", False)
        self.run(p, ")")
        self.stats["display_equations"] += 1
        self.stats["paragraphs"] += 1

    # -------------------------------------------------------------- blocks
    def para(self, text, size=BODY_PT, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, before=0, after=4,
             spacing=1.5, style=None, ea=EA_BODY, keep_next=False, hanging=None):
        p = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), spacing
        if align is not None:
            p.alignment = align
        if indent:
            pf.first_line_indent = Pt(2 * size)
        if hanging:
            pf.left_indent, pf.first_line_indent = Pt(hanging), Pt(-hanging)
        pf.keep_with_next = keep_next
        self.rich(p, text, size=size, bold=bold, ea=ea)
        self.stats["paragraphs"] += 1
        return p

    def reference(self, s):
        n, body = REF_RX.match(s).groups()
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent = Pt(28), Pt(-28)
        pf.space_after, pf.line_spacing = Pt(2), 1.2
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self.bookmark(p, f"ref_{int(n)}", True)
        self.run(p, f"[{n}]", SMALL_PT)
        self.bookmark(p, "", False)
        self.run(p, "\t", SMALL_PT)
        m = LINK_RX.search(body)
        if m:
            self.plain(p, body[:m.start()], SMALL_PT, False, False, EA_BODY)
            url = f"https://doi.org/{m.group(2)}" if m.group(2) else m.group(3)
            self.link(p, m.group(1), SMALL_PT, url=url)
            self.run(p, ".", SMALL_PT)
        else:
            self.plain(p, body, SMALL_PT, False, False, EA_BODY)
        self.stats["references"] += 1
        self.stats["paragraphs"] += 1

    def table(self, rows):
        ncol = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False                                   # fixed layout with content-based column widths
        for j, w in enumerate(column_widths(rows, ncol)):
            t.columns[j].width = Twips(w)
            for cell in t.columns[j].cells:
                cell.width = Twips(w)
        centre = [j > 0 and numeric_column([r[j] if j < len(r) else "" for r in rows]) for j in range(ncol)]
        for i, row in enumerate(rows):
            for j in range(ncol):
                cell = t.cell(i, j)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.1
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centre[j] else WD_ALIGN_PARAGRAPH.LEFT
                self.rich(p, row[j] if j < len(row) else "", size=TABLE_PT, bold=(i == 0))
                set_cell_border(cell, top=(i == 0), bottom=(i in (0, len(rows) - 1)))
        self.stats["tables"] += 1

    def figure(self, stem, images, new_page=True):
        if not images:
            return
        src = FIG / f"{stem}.png"
        if not src.exists():
            self.stats["missing_figures"].append(str(src))
            return
        self.doc.add_picture(str(review_png(stem)), width=Cm(16.0))
        pf = self.doc.paragraphs[-1].paragraph_format
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.keep_with_next = True
        pf.page_break_before = new_page                          # "page break before" never leaves an empty page
        self.stats["images"] += 1

    # -------------------------------------------------------------- Markdown -> document
    def build(self, md_path: Path, out_path: Path, images: bool) -> dict:
        lines = md_path.read_text(encoding="utf-8").lstrip("﻿").split("\n")
        kind = self.kind
        section, first_table_in_section, i = "", True, 0
        after_heading = False
        while i < len(lines):
            s = lines[i].strip()
            if not s:
                i += 1
                continue
            m = FIG_RX.match(s)
            if m:
                self.figure(m.group(1), images, new_page=not after_heading)
                after_heading = False
                i += 1
                continue
            after_heading = bool(re.match(r"^#{1,4}\s", s))
            if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
                block = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    block.append(lines[i])
                    i += 1
                self.table(parse_markdown_table(block))
                self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
                continue
            m = re.match(r"^(#{1,4})\s+(.*)$", s)
            if m:
                level, title = len(m.group(1)), m.group(2)
                section = title
                style = {1: "Title", 2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}[level]
                p = self.doc.add_paragraph(style=style)
                p.paragraph_format.page_break_before = (kind == "main" and title in ("参考文献", "表", "图注")) or \
                    (kind == "supplement" and title in ("补充表", "补充图", "补充参考文献"))
                self.rich(p, title, size={1: 16, 2: 14, 3: 12, 4: 12}[level], bold=True, ea=EA_HEAD)
                self.stats["headings"] += 1
                first_table_in_section = True
                i += 1
                continue
            m = MT.DISPLAY_RX.match(s)
            if m:
                self.display(m.group(1).strip())
                i += 1
                continue
            if section in ("参考文献", "补充参考文献") and REF_RX.match(s):
                self.reference(s)
                i += 1
                continue
            if re.match(r"^[-*]\s+", s):
                p = self.doc.add_paragraph(style="List Bullet")
                p.paragraph_format.line_spacing, p.paragraph_format.space_after = 1.5, Pt(2)
                self.rich(p, re.sub(r"^[-*]\s+", "", s))
                self.stats["bullets"] += 1
                i += 1
                continue
            if CAPTION_RX.match(s):
                brk = False
                if s.startswith("表") and section in ("表", "补充表"):
                    brk = not first_table_in_section and kind != "supplement"
                    first_table_in_section = False
                head = s.startswith("表")
                p = self.para(s, size=SMALL_PT, bold=head, indent=False, before=8 if head else 4, after=4, keep_next=True,
                              style="Caption",
                              align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.2)
                p.paragraph_format.page_break_before = brk
                i += 1
                continue
            if s.startswith("注："):
                self.para(s, size=TABLE_PT, indent=False, after=6, spacing=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
                i += 1
                continue
            if s.startswith("**图 "):
                self.para(s, size=SMALL_PT, indent=False, before=4, after=10, spacing=1.3, style="Caption")
                i += 1
                continue
            if s.startswith(("**短标题**", "**作者**", "¹", "**关键词**")) or \
                    (kind == "supplement" and s.startswith("**") and s.endswith("**")):
                self.para(s, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, after=6)
                i += 1
                continue
            self.para(s)
            i += 1
        if self.n_obj:
            pkg = self.doc.part.package
            pkg.relate_to(Part(PackURI("/docProps/custom.xml"), CT.OFC_CUSTOM_PROPERTIES, CUSTOM_XML.encode("utf-8"), pkg),
                          RT.CUSTOM_PROPERTIES)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out_path))
        return audit(out_path, self.stats)


# ------------------------------------------------------------------ audit of a saved document
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W, "o": "urn:schemas-microsoft-com:office:office", "v": "urn:schemas-microsoft-com:vml",
      "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}


def audit(path: Path, stats: dict) -> dict:
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        root = etree.fromstring(z.read("word/document.xml"))
        rels = {r.get("Id"): r for r in etree.fromstring(z.read("word/_rels/document.xml.rels")).findall("rel:Relationship", NS)}
        ctypes = z.read("[Content_Types].xml").decode("utf-8")
        custom = z.read("docProps/custom.xml").decode("utf-8") if "docProps/custom.xml" in names else ""
        oles = root.findall(".//o:OLEObject", NS)
        bins = []
        for o in oles:
            target = rels[o.get(f"{{{NS['r']}}}id")].get("Target")
            bins.append(MT.audit_bin(z.read("word/" + target)))
        imgs = [rels[s.get(f"{{{NS['r']}}}id")].get("Target") for s in root.findall(".//v:imagedata", NS)]
        img_ok = all(("word/" + t) in names for t in imgs)
    body_text = "".join(t.text or "" for t in root.iter(f"{{{W}}}t"))
    instr = [t.text.strip() for t in root.iter(f"{{{W}}}instrText")]
    bms = {b.get(f"{{{W}}}name") for b in root.iter(f"{{{W}}}bookmarkStart")}
    anchors = [h.get(f"{{{W}}}anchor") for h in root.iter(f"{{{W}}}hyperlink") if h.get(f"{{{W}}}anchor")]
    disp = [p for p in root.iter(f"{{{W}}}p")
            if (s := p.find("w:pPr/w:pStyle", NS)) is not None and s.get(f"{{{W}}}val") == "MTDisplayEquation"]
    disp_ok = all(len(p.findall(".//w:tab", NS)) == 2 and len(p.findall(".//w:object", NS)) == 1 for p in disp)
    seq = [x for x in instr if x.startswith("SEQ MTEqn")]
    refs = [x.split()[1] for x in instr if x.startswith("REF ")]
    d = Document(str(path))
    sec = d.sections[0]
    stats.update({
        "path": str(path), "bytes": path.stat().st_size,
        "a4": abs(sec.page_width.cm - 21.0) < 0.05 and abs(sec.page_height.cm - 29.7) < 0.05,
        "line_numbers": sec._sectPr.find(qn("w:lnNumType")) is not None,
        "page_field": any("PAGE" in (z_.text or "") for f in d.sections for p in f.footer.paragraphs
                          for z_ in p._p.iter(qn("w:instrText"))),
        "omml_objects": len(root.findall(".//m:oMath", NS)),
        "ole_objects": len(oles), "ole_progid_ok": all(o.get("ProgID") == "Equation.DSMT4" for o in oles),
        "ole_bins_ok": all(b["mtef_ok"] and b["replacement_chars"] == 0 for b in bins),
        "ole_embeddings": sum(1 for n in names if n.startswith("word/embeddings/")),
        "previews_ok": img_ok and len(imgs) == len(oles),
        "content_types_ok": ("image/x-wmf" in ctypes) and ("oleObject" in ctypes or not oles),
        "mtwineqns": ("MTWinEqns" in custom) or not oles,
        "display_paragraphs": len(disp), "display_layout_ok": disp_ok, "seq_fields": len(seq),
        "eq_numbers_ok": [int(re.search(r"(\d+)$", b).group(1)) - stats_base(stats) for b in sorted(
            x for x in bms if x and x.startswith("ZEqnNum"))] == list(range(1, len(disp) + 1)),
        "eq_refs_resolve": all(r in bms for r in refs), "n_eq_refs": len(refs),
        "citation_anchors_resolve": all(a in bms for a in anchors), "n_citation_links": len(anchors),
        "verified_tables": len(d.tables), "verified_images": sum(1 for n in names if n.startswith("word/media/")
                                                                  and not n.endswith(".wmf")),
        "heading_styles": sum(1 for p in d.paragraphs if p.style.name.startswith(("Heading", "Title"))),
        "leftover_dollar": "$" in body_text, "leftover_markup": bool(re.search(r"<sup>|\*\*|@eq:|<<", body_text)),
    })
    stats["cited_numbers"] = sorted(set(stats["cited_numbers"]))
    stats["ok"] = bool(
        stats["a4"] and stats["line_numbers"] == stats["want_line_numbers"] and stats["page_field"]
        and stats["omml_objects"] == 0 and stats["ole_objects"] == stats["objects"] == stats["ole_embeddings"]
        and stats["ole_progid_ok"] and stats["ole_bins_ok"] and stats["previews_ok"] and stats["content_types_ok"]
        and stats["mtwineqns"] and stats["display_paragraphs"] == stats["display_equations"] == stats["seq_fields"]
        and stats["display_layout_ok"] and stats["eq_numbers_ok"] and stats["eq_refs_resolve"]
        and stats["citation_anchors_resolve"] and not stats["missing_figures"] and not stats["missing_assets"]
        and not stats["leftover_dollar"] and not stats["leftover_markup"])
    return stats


def stats_base(stats: dict) -> int:
    return 900000 if stats["kind"] == "supplement" else 800000


def load_assets(md_paths) -> dict:
    """{(tex, mode): info} for every equation of the given Markdown files (all must have been converted)."""
    out = {}
    for tex, mode in MT.collect(md_paths):
        j = MT.CACHE / f"{MT.eq_key(tex, mode)}.json"
        if j.exists():
            out[(tex, mode)] = json.loads(j.read_text(encoding="utf-8"))
    return out


def run() -> dict:
    main_md, supp_md, notes_md = ROOT / f"{STEM}.md", ROOT / f"{SUPP_STEM}.md", ROOT / f"{NOTES_STEM}.md"
    mds = [p for p in (main_md, supp_md, notes_md) if p.exists()]
    assets = load_assets(mds)
    rep = {"date": DATE, "assets": len(assets)}
    rep["submission"] = Builder("main", assets).build(main_md, ROOT / f"{STEM}.docx", images=False)
    rep["review"] = Builder("main", assets).build(main_md, ROOT / f"{REVIEW_STEM}.docx", images=True)
    rep["supplement"] = Builder("supplement", assets).build(supp_md, ROOT / f"{SUPP_STEM}.docx", images=True)
    if notes_md.exists():
        rep["notes"] = Builder("notes", assets, line_numbers=False).build(notes_md, notes_md.with_suffix(".docx"),
                                                                           images=False)
    rep["all_ok"] = all(rep[k]["ok"] for k in ("submission", "review", "supplement", "notes") if k in rep)
    (ROOT / "build").mkdir(exist_ok=True)
    (ROOT / "build" / "docx_build_report10.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1),
                                                             encoding="utf-8")
    return rep


if __name__ == "__main__":
    r = run()
    print(json.dumps({k: ({kk: vv for kk, vv in v.items() if kk not in ("cited_numbers", "eq_refs")}
                          if isinstance(v, dict) else v) for k, v in r.items()}, ensure_ascii=False, indent=1))
