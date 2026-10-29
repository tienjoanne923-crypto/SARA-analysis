"""Word documents of the English SARA submission (Briefings in Bioinformatics style).

Built from the Markdown written by ``manuscript_en`` / ``supp_en`` with the machinery of ``build_docx10``: every
equation is a MathType object (Equation.DSMT4 OLE storage and WMF preview made by MathType itself, size-scaled for 9 pt
tables and 10.5 pt legends), display equations use MathType's MTDisplayEquation layout with ``SEQ MTEqn`` numbers in
``ZEqnNum`` bookmarks, and every "Eq. (n)" in the text is a ``REF`` field to that bookmark.  English rules:
  * citations are bracketed numbers in the text line ([1], [2, 5], [1–4]); every number is hyperlinked to the
    bookmark of its reference entry, and the reference list is numbered "1." with a hanging indent;
  * captions "Table n." (label in bold) above tables, legends "**Figure n. Title.**", table notes "Note:";
  * Times New Roman for every script slot, document language en-US, left-aligned paragraphs without indent.
Every saved file is reopened and audited (``build_docx10.audit``); the report goes to build/en/docx_build_report_en.json.
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips

from . import mt_factory10 as MT
from .build_docx9 import BODY_PT, FIG_RX, LATIN, SMALL_PT, TABLE_PT, parse_markdown_table, set_cell_border, style_fonts
from .build_docx10 import (CUSTOM_XML, ESC, LINK_RX, PAD, PVAL, TEXT_W, UNIT, Builder, _cw, _plain, _units, audit,
                           load_assets)
from .manuscript_en import DATE, EN, STEM, TITLE
from .numbers9 import ROOT
from .supp_en import SUPP_STEM

REVIEW_STEM = f"SARA_manuscript_review_{DATE}"
FIG_EN = EN / "figures"
BUILD = ROOT / "build" / "en"
TOKEN = re.compile(r"(\$[^$]+\$|\[\d+(?:(?:–|, )\d+)*\]|\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+?`|Eq\. \(S?\d+\))")
CAPTION_RX = re.compile(r"^(Table S?\d+\.|\([a-z]\)) ")
REF_RX = re.compile(r"^(\d+)\.\s+(.*)$")
FRONT = ("**Running title", "[Author name]", "¹", "\\*Corresponding", "**Keywords")
BREAKS = {"main": ("References", "Tables", "Figure legends"),
          "supplement": ("Supplementary tables", "Supplementary figures", "Supplementary references")}
REF_SECTIONS = ("References", "Supplementary references")
NUMERIC_CELL = re.compile(r"[+−\-]?[\d.,]+%?( \([^)]*\))?[†‡§]?"                # a value, optionally with its interval
                          r"|[+−\-]?[\d.,]+%?( ?[/±] ?[+−\-]?[\d.,]+%?)+"          # a / b (/ c), a ± b
                          r"|\d{4}-\d{2}(-\d{2})?|—|Reference|Yes|No")


def numeric_column(col: list[str]) -> bool:
    cells = [c.strip() for c in col[1:] if c.strip()]
    return bool(cells) and sum(bool(NUMERIC_CELL.fullmatch(c.replace("**", ""))) for c in cells) >= 0.6 * len(cells)


def _token(s: str) -> float:
    """widest unbreakable piece of English cell text (Word breaks at spaces and after hyphens, en dashes and slashes),
    in half-em units at 9 pt."""
    pieces = [t for t in re.split(r"\s+|(?<=[-–/])", _plain(s)) if t]
    return max((sum(_cw(c) for c in t) for t in pieces), default=0) + 1


def column_widths(rows, ncol: int, total: int = TEXT_W) -> list[int]:
    """twips per column (build_docx10.column_widths with English break rules): every column keeps its widest
    unbreakable piece; if the table fits, nothing wraps; otherwise headers may wrap to two lines and the widest columns
    wrap to a common cap.  The table never exceeds the text width."""
    cells = [[r[j] if j < len(r) else "" for r in rows] for j in range(ncol)]
    body = [max((_units(c) for c in col[1:]), default=0) for col in cells]
    head = [_units(col[0]) * 1.08 for col in cells]
    floor = [max(_token(c) * (1.08 if i == 0 else 1) for i, c in enumerate(col)) * UNIT + PAD for col in cells]
    if sum(floor) > total:
        floor = [f * total / sum(floor) for f in floor]
    # values with their intervals stay on one line whenever the text columns can absorb the wrapping
    num_w = [max((_units(c) for c in col[1:] if NUMERIC_CELL.fullmatch(c.strip().replace("**", ""))), default=0)
             for col in cells]
    keep = [max(f, b * UNIT + PAD) if j > 0 and numeric_column(col) else f
            for j, (f, b, col) in enumerate(zip(floor, num_w, cells))]
    if sum(keep) <= total:
        floor = keep
    full = [max(b, h, 4) * UNIT + PAD for b, h in zip(body, head)]
    if sum(full) <= total:
        return [int(w * total / sum(full)) for w in full]
    nat = [max(max(b, h / 2, 4) * UNIT + PAD, f) for b, h, f in zip(body, head, floor)]
    if sum(nat) <= total:
        return [int(w * total / sum(nat)) for w in nat]
    lo, hi = 0.0, max(nat)
    for _ in range(60):
        c = (lo + hi) / 2
        lo, hi = (c, hi) if sum(max(min(w, c), f) for w, f in zip(nat, floor)) <= total else (lo, c)
    return [int(max(min(w, lo), f)) for w, f in zip(nat, floor)]


def review_png(stem: str) -> Path:
    """300 dpi copy of the 600 dpi English figure (review documents only)."""
    from PIL import Image
    out = BUILD / "png300" / f"{stem}.png"
    src = FIG_EN / f"{stem}.png"
    if not out.exists() or out.stat().st_mtime < src.stat().st_mtime:
        out.parent.mkdir(parents=True, exist_ok=True)
        im = Image.open(src).convert("RGB")
        im = im.resize((im.width // 2, im.height // 2), Image.LANCZOS)
        im.save(out, dpi=(300, 300), optimize=True)
    return out


class EnBuilder(Builder):
    def __init__(self, kind: str, assets: dict, line_numbers: bool = True):
        super().__init__(kind, assets, line_numbers=line_numbers)
        st = self.doc.styles
        style_fonts(st["Normal"], LATIN, BODY_PT)
        for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
            style_fonts(st[name], LATIN, size, bold=True)
        style_fonts(st["List Bullet"], LATIN, BODY_PT)
        rpr = st.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        for k in ("w:val", "w:eastAsia"):
            lang.set(qn(k), "en-US")
        for s in self.doc.sections:
            for r in s.footer.paragraphs[0].runs:
                r._element.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:eastAsia"), LATIN)
        cp = self.doc.core_properties
        cp.title, cp.language = TITLE, "en-US"
        self.stats["cited_expanded"] = []

    # -------------------------------------------------------------- runs: Latin font in every script slot
    def run(self, p, text, size=BODY_PT, bold=False, italic=False, ea=None, mono=False, sup=False):
        return super().run(p, text, size, bold, italic, LATIN, mono, sup)

    def rich(self, p, text, size=BODY_PT, bold=False, italic=False, ea=None):
        for tok in TOKEN.split((text or "").replace("\\*", ESC)):
            if not tok:
                continue
            if tok.startswith("$") and tok.endswith("$") and len(tok) > 2:
                self.inline_math(p, tok[1:-1], size)
            elif re.fullmatch(r"\[\d+(?:(?:–|, )\d+)*\]", tok):
                self.citation(p, tok[1:-1], size, bold, LATIN)
            elif tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                self.rich(p, tok[2:-2], size, True, italic)
            elif tok.startswith("`") and tok.endswith("`"):
                self.run(p, tok[1:-1], size - 1, bold, italic, mono=True)
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                self.rich(p, tok[1:-1], size, bold, True)
            elif tok.startswith("Eq. ("):
                self.eq_ref(p, tok, size, bold, LATIN)
            else:
                self.plain(p, tok, size, bold, italic, LATIN)
        return p

    # -------------------------------------------------------------- citations and equation references
    def citation(self, p, inner, size, bold, ea):
        self.run(p, "[", size, bold)
        for k, part in enumerate(inner.split(", ")):
            if k:
                self.run(p, ", ", size, bold)
            ends = [int(x) for x in part.split("–")]
            for m, n in enumerate(ends):
                if m:
                    self.run(p, "–", size, bold)
                self.link(p, str(n), size, bold, LATIN, anchor=f"ref_{n}")
                self.stats["cited_numbers"].append(n)
            self.stats["cited_expanded"] += list(range(ends[0], ends[-1] + 1))
        self.run(p, "]", size, bold)
        self.stats["citations"] += 1

    def eq_ref(self, p, tok, size, bold, ea):
        m = re.fullmatch(r"Eq\. \((S?)(\d+)\)", tok)
        n = int(m.group(2))
        self.run(p, "Eq. (" + m.group(1), size, bold)
        self.field(p, f"REF {self.eq_bm(n)} \\h \\* MERGEFORMAT", str(n), size, bold, LATIN)
        self.run(p, ")", size, bold)
        self.stats["eq_refs"].append(n)

    # -------------------------------------------------------------- blocks
    def para(self, text, size=BODY_PT, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, before=0, after=6,
             spacing=1.5, style=None, ea=None, keep_next=False, hanging=None):
        return super().para(text, size, bold, align, indent, before, after, spacing, style, LATIN, keep_next, hanging)

    def reference(self, s):
        n, body = REF_RX.match(s).groups()
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent = Pt(24), Pt(-24)
        pf.space_after, pf.line_spacing = Pt(3), 1.2
        pf.tab_stops.add_tab_stop(Pt(24))
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self.bookmark(p, f"ref_{int(n)}", True)
        self.run(p, f"{n}.", SMALL_PT)
        self.bookmark(p, "", False)
        self.run(p, "\t", SMALL_PT)
        m = LINK_RX.search(body)
        if m:
            self.plain(p, body[:m.start()], SMALL_PT, False, False, LATIN)
            url = f"https://doi.org/{m.group(2)}" if m.group(2) else m.group(3)
            self.link(p, m.group(1), SMALL_PT, url=url)
            self.run(p, ".", SMALL_PT)
        else:
            self.plain(p, body, SMALL_PT, False, False, LATIN)
        self.stats["references"] += 1
        self.stats["paragraphs"] += 1

    def table(self, rows):
        ncol = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for j, w in enumerate(column_widths(rows, ncol)):
            t.columns[j].width = Twips(w)
            for cell in t.columns[j].cells:
                cell.width = Twips(w)
        centre = [j > 0 and numeric_column([r[j] if j < len(r) else "" for r in rows]) for j in range(ncol)]
        for i, row in enumerate(rows):
            for j in range(ncol):
                cell = t.cell(i, j)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.1
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centre[j] else WD_ALIGN_PARAGRAPH.LEFT
                self.rich(p, row[j] if j < len(row) else "", size=TABLE_PT, bold=(i == 0))
                set_cell_border(cell, top=(i == 0), bottom=(i in (0, len(rows) - 1)))
        self.stats["tables"] += 1

    def figure(self, stem, images, new_page=True):
        if not images:
            return
        src = FIG_EN / f"{stem}.png"
        if not src.exists():
            self.stats["missing_figures"].append(str(src))
            return
        self.doc.add_picture(str(review_png(stem)), width=Cm(16.0))
        para = self.doc.paragraphs[-1]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.page_break_before = new_page
        self.stats["images"] += 1

    # -------------------------------------------------------------- Markdown -> document
    def build(self, md_path: Path, out_path: Path, images: bool) -> dict:
        lines = md_path.read_text(encoding="utf-8").lstrip("\ufeff").split("\n")
        kind = self.kind
        section, first_table_in_section, i, after_heading = "", True, 0, False
        while i < len(lines):
            s = lines[i].strip()
            if not s:
                i += 1
                continue
            m = FIG_RX.match(s)
            if m:
                self.figure(m.group(1), images, new_page=not after_heading)
                after_heading = False
                i += 1
                continue
            after_heading = bool(re.match(r"^#{1,4}\s", s))
            if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
                block = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    block.append(lines[i])
                    i += 1
                self.table(parse_markdown_table(block))
                self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
                continue
            m = re.match(r"^(#{1,4})\s+(.*)$", s)
            if m:
                level, title = len(m.group(1)), m.group(2)
                section = title
                style = {1: "Title", 2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}[level]
                p = self.doc.add_paragraph(style=style)
                p.paragraph_format.page_break_before = title in BREAKS.get(kind, ())
                self.rich(p, title, size={1: 16, 2: 14, 3: 12, 4: 12}[level], bold=True)
                self.stats["headings"] += 1
                first_table_in_section = True
                i += 1
                continue
            m = MT.DISPLAY_RX.match(s)
            if m:
                self.display(m.group(1).strip())
                i += 1
                continue
            if section in REF_SECTIONS and REF_RX.match(s):
                self.reference(s)
                i += 1
                continue
            if re.match(r"^[-*]\s+", s):
                p = self.doc.add_paragraph(style="List Bullet")
                p.paragraph_format.line_spacing, p.paragraph_format.space_after = 1.5, Pt(3)
                self.rich(p, re.sub(r"^[-*]\s+", "", s))
                self.stats["bullets"] += 1
                i += 1
                continue
            m = CAPTION_RX.match(s)
            if m:
                head = s.startswith("Table")
                brk = False
                if head and section in ("Tables", "Supplementary tables"):
                    brk = not first_table_in_section and kind != "supplement"
                    first_table_in_section = False
                text = f"**{m.group(1)}** {s[m.end():]}" if head else s
                p = self.para(text, size=SMALL_PT, before=8 if head else 4, after=4, keep_next=True, spacing=1.2)
                p.paragraph_format.page_break_before = brk
                i += 1
                continue
            if s.startswith("Note:"):
                self.para(s, size=TABLE_PT, after=6, spacing=1.2)
                i += 1
                continue
            if s.startswith("**Figure "):
                self.para(s, size=SMALL_PT, before=4, after=10, spacing=1.3)
                i += 1
                continue
            if s.startswith(FRONT) or (kind == "supplement" and s.startswith("**") and s.endswith("**")):
                self.para(s, after=6)
                i += 1
                continue
            self.para(s)
            i += 1
        if self.n_obj:
            pkg = self.doc.part.package
            pkg.relate_to(Part(PackURI("/docProps/custom.xml"), CT.OFC_CUSTOM_PROPERTIES, CUSTOM_XML.encode("utf-8"), pkg),
                          RT.CUSTOM_PROPERTIES)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out_path))
        st = audit(out_path, self.stats)
        st["cited_expanded"] = sorted(set(st["cited_expanded"]))
        return st


# ------------------------------------------------------------------ documents, size-specific MathType objects, PDFs
def mds() -> list[Path]:
    return [EN / f"{STEM}.md", EN / f"{SUPP_STEM}.md"]


def requests() -> tuple[dict, set]:
    """(assets, {(tex, mode, size)}) of every object placed in the two English documents."""
    paths = mds()
    assets = load_assets(paths)
    req = set()
    (MT.CACHE / "work").mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=MT.CACHE / "work") as tmp:
        for md, kind in zip(paths, ("main", "supplement")):
            b = EnBuilder(kind, assets)
            b.build(md, Path(tmp) / f"{kind}.docx", images=False)
            req |= set(b.requests)
    return assets, req


def run() -> dict:
    main_md, supp_md = mds()
    assets = load_assets([main_md, supp_md])
    rep = {"date": DATE, "assets": len(assets)}
    rep["submission"] = EnBuilder("main", assets).build(main_md, EN / f"{STEM}.docx", images=False)
    rep["review"] = EnBuilder("main", assets).build(main_md, EN / f"{REVIEW_STEM}.docx", images=True)
    rep["supplement"] = EnBuilder("supplement", assets).build(supp_md, EN / f"{SUPP_STEM}.docx", images=True)
    rep["all_ok"] = all(rep[k]["ok"] for k in ("submission", "review", "supplement"))
    BUILD.mkdir(parents=True, exist_ok=True)
    (BUILD / "docx_build_report_en.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    return rep


def pdfs(render: bool = True) -> dict:
    """PDFs exported by Word itself (hidden instance on a copy); the submission PDF is for inspection only."""
    from .wordpdf10 import export
    qa = BUILD / "pdf_word"
    qa.mkdir(parents=True, exist_ok=True)
    jobs = {"submission": (EN / f"{STEM}.docx", qa / "submission_word.pdf"),
            "review": (EN / f"{REVIEW_STEM}.docx", EN / f"{REVIEW_STEM}.pdf"),
            "supplement": (EN / f"{SUPP_STEM}.docx", EN / f"{SUPP_STEM}.pdf")}
    rep = {}
    for name, (docx, pdf) in jobs.items():
        rep[name] = export(docx, pdf)
        if render:
            import shutil
            sys.path.append(r"D:\信\gpu8\tools_libs10")
            from .qa10 import render as png
            src = pdf if pdf.parent == qa else Path(shutil.copy2(pdf, qa / pdf.name))
            out = png(src, dpi=90)
            rep[name]["png_dir"], rep[name]["blank_pages"] = out["png_dir"], out["blank_pages"]
    (BUILD / "word_pdf_report_en.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    return rep


def roundtrip() -> dict:
    """MathType re-typesets every object of copies of the English submission and supplement files from the stored MTEF;
    each object must come back with the size written here."""
    from .mt_roundtrip10 import SCRIPT, objects
    rep = {}
    work = MT.CACHE / "work"
    for name, src in (("submission", EN / f"{STEM}.docx"), ("supplement", EN / f"{SUPP_STEM}.docx")):
        conv = work / f"roundtrip_en_{name}_mt.docx"
        conv.unlink(missing_ok=True)
        r = subprocess.run(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(SCRIPT), "-InputPath",
                            str(src), "-OutputPath", str(conv), "-TimeoutSeconds", "1200", "-SourceMathType"],
                           capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=1500)
        d = {"returncode": r.returncode, "stdout": r.stdout[-600:]}
        if r.returncode == 0:
            a, b = objects(src), objects(conv)
            ratios = [y["dya"] / x["dya"] for x, y in zip(a, b)] + [y["dxa"] / x["dxa"] for x, y in zip(a, b)]
            d.update({"objects": len(a), "objects_after": len(b), "identical_size": sum(
                1 for x, y in zip(a, b) if (x["dxa"], x["dya"]) == (y["dxa"], y["dya"])),
                "ratio_min": round(min(ratios), 3), "ratio_max": round(max(ratios), 3)})
            d["ok"] = len(a) == len(b) and 0.97 <= d["ratio_min"] and d["ratio_max"] <= 1.03
        else:
            d["ok"] = False
        rep[name] = d
    rep["ok"] = all(v["ok"] for v in rep.values() if isinstance(v, dict))
    (BUILD / "roundtrip_en.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    return rep


if __name__ == "__main__":
    arg = sys.argv[1] if len(sys.argv) > 1 else "build"
    if arg == "build":
        r = run()
        print(json.dumps({k: ({kk: vv for kk, vv in v.items() if kk not in ("cited_numbers", "eq_refs", "cited_expanded")}
                              if isinstance(v, dict) else v) for k, v in r.items()}, ensure_ascii=False, indent=1))
    elif arg == "requests":
        a, req = requests()
        todo = sorted({(t, m, s) for t, m, s in req if abs(s / 12 - 1) > 1e-9 and not (
            MT.CACHE / f"{a[(t, m)]['key']}_{round(s / 12 * 1000)}_rt.json").exists()})
        print(json.dumps({"placed": len(req), "scaled": sum(1 for *_, s in req if s != 12), "missing_rt": len(todo),
                          "todo": todo}, ensure_ascii=False, indent=1))
    elif arg == "pdf":
        print(json.dumps(pdfs(), ensure_ascii=False, indent=1))
    elif arg == "roundtrip":
        print(json.dumps(roundtrip(), ensure_ascii=False, indent=1))
